"""المشروع الحر (free_project + web/routes_free): إنشاء، تعديل البنية، استيراد، رفع الوسائط، والحدود مع المهمة الموجّهة."""

import io
import json
import os
import shutil
import sys
from pathlib import Path

import openpyxl
import pytest
from fastapi import FastAPI
from fastapi.testclient import TestClient
from PIL import Image

sys.path.insert(0, os.path.dirname(__file__))
from a_helpers import env, make_script  # noqa: E402,F401
from test_free_media import make_video  # noqa: E402

import config  # noqa: E402
import free_project as fp  # noqa: E402
import image_studio  # noqa: E402
import scene_import  # noqa: E402
import studio_gate  # noqa: E402
import timeline  # noqa: E402

FFMPEG = shutil.which("ffmpeg") is not None


def script_of(path):
    return json.loads(Path(path).read_text(encoding="utf-8"))


def img_file(tmp_path, name="a.png", w=320, h=480, color=(10, 120, 200)):
    p = tmp_path / name
    Image.new("RGB", (w, h), color).save(p)
    return p


def sp(made):
    return Path(config.SCRIPTS_DIR) / made["file"]


def rows(*texts):
    return [{"voiceover": t, "visual": "", "on_screen_text": "", "duration_sec": 3.0} for t in texts]


def struct(board, mutate=None):
    """شكل payload لـ apply_structure من لوحة موجودة (بنفس المعرّفات)."""
    s = {"rev": board["rev"], "scenes": [
        {"key": sc["key"], "voiceover": sc["voiceover"], "visual": sc["visual"], "emotion": sc["emotion"],
         "shots": [{"id": sh["id"], "duration_ms": sh["end_ms"] - sh["start_ms"], "text": sh["text"]} for sh in sc["shots"]]}
        for sc in board["scenes"]]}
    if mutate:
        mutate(s)
    return s


# ───────────────────────── الإنشاء ─────────────────────────
def test_blank_project_has_one_scene_one_shot_and_mirrors_a_free_script(env):
    made = fp.create("مشروعي")
    p = sp(made)
    s = script_of(p)
    assert s["_meta"]["mode"] == "free" and s["_meta"]["source"] == "blank" and s["_meta"]["display_name"] == "✎ مشروعي"
    assert s["_meta"]["script_id"].startswith("FREE-") and s["title"] == "مشروعي" and len(s["scenes"]) == 1
    tl = timeline.load(p.stem)
    assert tl["mode"] == "free" and tl["rev"] == 1 and tl["length_ms"] == 3000 and timeline.validate(tl)["ok"]
    sc = tl["scenes"][0]
    assert sc["key"] == "c001" and sc["n"] == 1 and sc["shots"][0]["id"] == "f001" and sc["shots"][0]["still_motion"] == "zoom_in"
    assert studio_gate.is_free(p)


def test_two_projects_with_the_same_title_get_different_files_and_script_ids(env):
    a, b = fp.create("تجربة"), fp.create("تجربة")
    assert a["file"] != b["file"]
    ida, idb = (script_of(sp(x))["_meta"]["script_id"] for x in (a, b))
    assert ida != idb


def test_create_from_imported_scenes_sets_durations_text_and_types(env):
    imp = scene_import.parse_text("ليش عيونك تتعب؟\nشاشة: عيونك تعبانة\nمدة: 4\nلقطة: قريبة\n\nجرّبها أسبوع.")["scenes"]
    made = fp.create("من نص", imp, source="text")
    tl = timeline.load(Path(made["file"]).stem)
    assert [sc["voiceover"] for sc in tl["scenes"]] == ["ليش عيونك تتعب؟", "جرّبها أسبوع."]
    s0 = tl["scenes"][0]["shots"][0]
    assert s0["text"] == "عيونك تعبانة" and s0["type"] == "close" and s0["end_ms"] - s0["start_ms"] == 4000
    assert tl["length_ms"] == 4000 + tl["scenes"][1]["shots"][0]["end_ms"] - tl["scenes"][1]["shots"][0]["start_ms"]
    sc = script_of(sp(made))
    assert sc["_meta"]["source"] == "text" and sc["scenes"][0]["on_screen_text"] == "عيونك تعبانة" and sc["hook_line"] == "ليش عيونك تتعب؟"
    assert sc["total_duration_sec"] == round(tl["length_ms"] / 1000, 2)


def test_title_is_trimmed_defaulted_and_limited(env):
    assert script_of(sp(fp.create("   ")))["title"] == "مشروع حر"
    with pytest.raises(fp.FreeError):
        fp.create("ع" * (fp.MAX_TITLE + 1))
    with pytest.raises(fp.FreeError) as e:
        fp.create("x", rows(*["م"] * (timeline.FREE_MAX_SCENES + 1)))
    assert "الحد" in str(e.value)


def test_attaching_a_product_copies_its_name_and_rejects_unknown_files(env):
    env.product("kama.json", product_name="KAMA", product_code="kama")
    made = fp.create("مع منتج", product_file="kama.json")
    meta = script_of(sp(made))["_meta"]
    assert meta["product"] == "KAMA" and meta["product_file"] == "kama.json" and meta["product_code"] == "kama"
    with pytest.raises(fp.FreeError) as e:
        fp.create("x", product_file="nope.json")
    assert e.value.status == 404


# ───────────────────────── تعديل البنية ─────────────────────────
def test_apply_structure_edits_text_and_durations_keeping_ids_and_bumping_rev(env):
    made = fp.create("p", rows("أول", "تاني"))
    p = sp(made)
    b = fp.board(p)

    def mut(s):
        s["scenes"][0]["voiceover"] = "كلام جديد"
        s["scenes"][0]["shots"][0].update(duration_ms=5000, text="نص", type="wide", still_motion="zoom_out", trim_in_ms=1500, keep_audio=True)

    out = fp.apply_structure(p, struct(b, mut))
    assert out["rev"] == b["rev"] + 1 and out["length_ms"] == 5000 + 3000
    sh = out["scenes"][0]["shots"][0]
    assert sh["id"] == "f001" and (sh["text"], sh["type"], sh["still_motion"], sh["trim_in_ms"], sh["keep_audio"]) == ("نص", "wide", "zoom_out", 1500, True)
    assert out["scenes"][1]["start_ms"] == 5000
    s = script_of(p)
    assert s["scenes"][0]["voiceover"] == "كلام جديد" and s["scenes"][0]["duration_sec"] == 5.0 and s["scenes"][0]["shot_type"] == "wide"
    assert s["total_duration_sec"] == 8.0 and s["hook_line"] == "كلام جديد"


def test_reorder_add_and_remove_scenes_keep_media_with_their_shot(env, tmp_path):
    made = fp.create("p", rows("أول", "تاني", "تالت"))
    p = sp(made)
    stem = p.stem
    tl = timeline.load(stem)
    tl["scenes"][2]["shots"][0]["image"]["chosen"] = "img/c.png"
    timeline.save(stem, tl)
    b = fp.board(p)

    def mut(s):
        sc = s["scenes"]
        sc[0], sc[2] = sc[2], sc[0]                                         # بدّل مكان الأول والتالت
        sc.append({"voiceover": "جديد", "shots": [{"duration_ms": 2000, "text": ""}]})   # مشهد جديد بلا معرّفات
        del sc[1]                                                           # احذف التاني (key c002)

    out = fp.apply_structure(p, struct(b, mut))
    assert [s["voiceover"] for s in out["scenes"]] == ["تالت", "أول", "جديد"]
    assert [s["n"] for s in out["scenes"]] == [1, 2, 3]
    assert [s["key"] for s in out["scenes"]][:2] == ["c003", "c001"] and out["scenes"][2]["key"] == "c004"
    assert out["scenes"][0]["shots"][0]["image"]["chosen"] == "img/c.png"        # صورة اللقطة مشت معها
    assert out["scenes"][2]["shots"][0]["id"] == "f004" and out["length_ms"] == 3000 + 3000 + 2000


def test_a_shot_can_move_between_scenes_and_a_scene_can_have_several_shots(env):
    made = fp.create("p", rows("أول", "تاني"))
    p = sp(made)
    b = fp.board(p)

    def mut(s):
        moved = s["scenes"][1]["shots"].pop()
        s["scenes"][0]["shots"].append(moved)
        s["scenes"][1]["shots"].append({"duration_ms": 1000, "text": "جديدة"})

    out = fp.apply_structure(p, struct(b, mut))
    assert [sh["id"] for sh in out["scenes"][0]["shots"]] == ["f001", "f002"] and out["scenes"][1]["shots"][0]["id"] == "f003"
    assert out["length_ms"] == 3000 + 3000 + 1000


def test_stale_rev_and_unknown_or_duplicate_ids_are_rejected_without_writing(env):
    made = fp.create("p", rows("أول", "تاني"))
    p = sp(made)
    b = fp.board(p)
    fp.apply_structure(p, struct(b))
    with pytest.raises(fp.FreeError) as e:
        fp.apply_structure(p, struct(b))                                    # نفس rev القديمة
    assert e.value.status == 409 and "حدّث الصفحة" in str(e.value)
    before = p.read_text(encoding="utf-8")
    b = fp.board(p)
    for bad in (lambda s: s["scenes"][0].update(key="c999"), lambda s: s["scenes"][0]["shots"][0].update(id="f999"),
                lambda s: s["scenes"][1].update(key="c001"), lambda s: s["scenes"][1]["shots"][0].update(id="f001"),
                lambda s: s["scenes"][0].update(shots=[]), lambda s: s["scenes"].append("x"),
                lambda s: s["scenes"][0]["shots"][0].update(duration_ms=100), lambda s: s["scenes"][0]["shots"][0].update(duration_ms=61_000),
                lambda s: s["scenes"][0]["shots"][0].update(duration_ms="abc"), lambda s: s["scenes"][0]["shots"][0].update(still_motion="spin"),
                lambda s: s["scenes"][0].update(voiceover="ك" * 1001), lambda s: s["scenes"][0]["shots"][0].update(text="ك" * 121)):
        with pytest.raises(fp.FreeError):
            fp.apply_structure(p, struct(b, bad))
    assert p.read_text(encoding="utf-8") == before and fp.board(p)["rev"] == b["rev"]


def test_title_and_product_can_change_and_detach(env):
    env.product("kama.json", product_name="KAMA", product_code="kama")
    made = fp.create("قديم")
    p = sp(made)
    b = fp.board(p)
    out = fp.apply_structure(p, {**struct(b), "title": "جديد", "product_file": "kama.json"})
    assert out["title"] == "جديد" and out["product"] == "KAMA" and script_of(p)["_meta"]["display_name"] == "✎ جديد"
    out = fp.apply_structure(p, {**struct(out), "product_file": ""})
    assert out["product"] == "" and out["product_file"] == ""
    with pytest.raises(fp.FreeError):
        fp.apply_structure(p, {**struct(out), "product_file": "missing.json"})


def test_structure_is_refused_for_guided_and_missing_scripts(env):
    g = env.script_file("g.json", make_script())
    with pytest.raises(fp.FreeError) as e:
        fp.board(g)
    assert e.value.status == 409
    with pytest.raises(fp.FreeError) as e:
        fp.board(Path(config.SCRIPTS_DIR) / "none.json")
    assert e.value.status == 404


# ───────────────────────── الاستيراد ─────────────────────────
def test_import_appends_or_replaces_and_respects_the_scene_limit(env):
    made = fp.create("p", rows("أول"))
    p = sp(made)
    out = fp.import_scenes(p, rows("تاني", "تالت"))
    assert [s["voiceover"] for s in out["scenes"]] == ["أول", "تاني", "تالت"] and out["rev"] == 2
    out = fp.import_scenes(p, rows("بس هاد"), replace=True)
    assert [s["voiceover"] for s in out["scenes"]] == ["بس هاد"]
    with pytest.raises(fp.FreeError):
        fp.import_scenes(p, rows(*["م"] * timeline.FREE_MAX_SCENES))


# ───────────────────────── الوسائط ─────────────────────────
def batch(tmp_path, *names, kind="image"):
    out = []
    for i, n in enumerate(names):
        p = img_file(tmp_path, f"u{i}.png", color=(i * 40, 100, 200)) if kind == "image" else None
        out.append({"kind": kind, "tmp": p, "filename": n})
    return out


def test_images_become_scenes_with_a_chosen_1080_frame_and_replace_the_blank_default(env, tmp_path):
    made = fp.create("p")                                                   # مشهد فاضي افتراضي
    p = sp(made)
    res = fp.add_media(p, batch(tmp_path, "a.png", "b.png"))
    b = res["board"]
    assert [len(sc["shots"]) for sc in b["scenes"]] == [1, 1] and b["length_ms"] == 6000
    for sc in b["scenes"]:
        sh = sc["shots"][0]
        assert sh["end_ms"] - sh["start_ms"] == 3000 and sh["still_motion"] == "zoom_in"
        chosen = image_studio.abs_of(p.stem, sh["image"]["chosen"])
        assert chosen and Image.open(chosen).size == (1080, 1920)
        assert sh["image"]["path_used"] == "upload" and sh["image"]["candidates"][0]["provider"] == "upload"
    assert b["readiness"]["ready"] and not b["readiness"]["no_media"]


def test_media_is_appended_after_existing_content(env, tmp_path):
    made = fp.create("p", rows("كلام"))
    p = sp(made)
    b = fp.add_media(p, batch(tmp_path, "a.png"))["board"]
    assert [sc["voiceover"] for sc in b["scenes"]] == ["كلام", ""] and b["readiness"]["no_media"] == ["f001"]


def test_one_bad_file_rejects_the_whole_batch_and_writes_nothing(env, tmp_path):
    made = fp.create("p")
    p = sp(made)
    bad = tmp_path / "bad.png"
    bad.write_bytes(b"not an image")
    items = batch(tmp_path, "good.png") + [{"kind": "image", "tmp": bad, "filename": "bad.png"}]
    before = fp.board(p)
    with pytest.raises(fp.FreeError) as e:
        fp.add_media(p, items)
    assert "bad.png" in str(e.value)
    after = fp.board(p)
    assert after["rev"] == before["rev"] and len(after["scenes"]) == 1 and not (Path(config.LIBRARY_DIR) / "timelines" / p.stem / "img").exists()


def test_batch_and_scene_limits(env, tmp_path):
    made = fp.create("p")
    p = sp(made)
    with pytest.raises(fp.FreeError) as e:
        fp.add_media(p, [{"kind": "image", "tmp": tmp_path / "x.png", "filename": "x.png"}] * (fp.MAX_IMAGES_PER_BATCH + 1))
    assert str(fp.MAX_IMAGES_PER_BATCH) in str(e.value)
    fp.import_scenes(p, rows(*["م"] * (timeline.FREE_MAX_SCENES - 1)), replace=True)
    with pytest.raises(fp.FreeError):
        fp.add_media(p, batch(tmp_path, "a.png", "b.png"))


def test_upload_to_a_specific_shot_makes_it_the_chosen_image(env, tmp_path):
    made = fp.create("p", rows("أول", "تاني"))
    p = sp(made)
    res = fp.upload_to_shot(p, "f002", img_file(tmp_path, "mine.png"), "mine.png")
    assert res["shot"]["id"] == "f002" and res["shot"]["image"]["chosen"] and not fp.board(p)["scenes"][0]["shots"][0]["image"]["chosen"]
    assert fp.board(p)["readiness"]["no_media"] == ["f001"]
    with pytest.raises(fp.FreeError) as e:
        fp.upload_to_shot(p, "f999", img_file(tmp_path, "x.png"), "x.png")
    assert e.value.status == 404
    with pytest.raises(fp.FreeError):
        fp.upload_to_shot(p, "f001", img_file(tmp_path, "y.png"), "y.exe")
    bad = tmp_path / "bad.png"
    bad.write_bytes(b"zzz")
    with pytest.raises(fp.FreeError):
        fp.upload_to_shot(p, "f001", bad, "bad.png")


@pytest.mark.skipif(not FFMPEG, reason="ffmpeg غير متوفر")
def test_videos_become_scenes_with_their_own_length_a_poster_and_audio_flag(env, tmp_path):
    made = fp.create("p")
    p = sp(made)
    v1 = make_video(tmp_path / "v1.mp4", 2, audio=True)
    v2 = make_video(tmp_path / "v2.mp4", 1)
    b = fp.add_media(p, [{"kind": "video", "tmp": v1, "filename": "v1.mp4"}, {"kind": "video", "tmp": v2, "filename": "v2.mp4"}])["board"]
    s1, s2 = b["scenes"][0]["shots"][0], b["scenes"][1]["shots"][0]
    assert 1800 <= s1["end_ms"] - s1["start_ms"] <= 2300 and 800 <= s2["end_ms"] - s2["start_ms"] <= 1300 or s2["end_ms"] - s2["start_ms"] == 500
    assert s1["motion"]["clip"].endswith(".mp4") and s1["motion"]["duration_ms"] == s1["end_ms"] - s1["start_ms"] and s1["motion"]["model_id"] == "upload"
    assert s1["keep_audio"] is True and s2["keep_audio"] is False                     # الأول فيه صوت والتاني لا
    assert s1["image"]["chosen"] and Image.open(image_studio.abs_of(p.stem, s1["image"]["chosen"])).size == (1080, 1920)
    assert not v1.exists() and not v2.exists()                                        # الملفات المؤقتة انتقلت لمجلد المقاطع


@pytest.mark.skipif(not FFMPEG, reason="ffmpeg غير متوفر")
def test_video_into_a_shot_with_a_voiced_scene_keeps_the_shot_length(env, tmp_path):
    made = fp.create("p", rows("أول"))
    p = sp(made)
    tl = timeline.load(p.stem)
    tl["scenes"][0]["voice"] = {"duration_ms": 3000, "path": "voice/x.wav"}
    timeline.save(p.stem, tl)
    res = fp.upload_to_shot(p, "f001", make_video(tmp_path / "v.mp4", 2), "v.mp4")
    sh = fp.board(p)["scenes"][0]["shots"][0]
    assert sh["end_ms"] - sh["start_ms"] == 3000 and sh["motion"]["duration_ms"] == 3000 and any("أقصر" in w for w in res["warnings"])


@pytest.mark.skipif(not FFMPEG, reason="ffmpeg غير متوفر")
def test_video_into_a_silent_scene_takes_the_clip_length_and_shifts_the_rest(env, tmp_path):
    made = fp.create("p", rows("أول", "تاني"))
    p = sp(made)
    fp.upload_to_shot(p, "f001", make_video(tmp_path / "v.mp4", 2), "v.mp4")
    b = fp.board(p)
    d0 = b["scenes"][0]["shots"][0]["end_ms"]
    assert 1800 <= d0 <= 2300 and b["scenes"][1]["start_ms"] == d0 and b["length_ms"] == d0 + 3000


# ───────────────────────── تنبيهات ─────────────────────────
def test_advisories_flag_forbidden_claims_and_fast_speech_but_never_block(env):
    imp = [{"voiceover": "هالقطرة بتشفي كل شي وبتعالج العيون نهائياً", "visual": "", "on_screen_text": "", "duration_sec": 2.0},
           {"voiceover": "كلمة " * 14, "visual": "", "on_screen_text": "سته كلمات بنص الشاشة هاد طويل", "duration_sec": 2.0}]
    made = fp.create("p", imp)
    p = sp(made)
    notes = fp.board(p)["advisories"]
    kinds = {n["kind"] for n in notes}
    assert {"claim", "pace", "screen_text"} <= kinds
    assert all(isinstance(n["message"], str) and n["message"] for n in notes)
    assert fp.board(p)["validation"]["ok"]


def test_clean_text_has_no_advisories(env):
    made = fp.create("p", rows("مرحبا بكم"))
    assert fp.board(sp(made))["advisories"] == []


# ───────────────────────── الحدود مع المهمة الموجّهة ─────────────────────────
def test_free_scripts_stay_out_of_hooks_bank_and_the_writers_used_registry(env):
    import hooks_bank
    import stage3_generate
    fp.create("p", rows("هوك حر مميّز جداً"))
    g = make_script(hook="هوك المهمة الموجّهة")
    env.script_file("g.json", g)
    texts = [e["text"] for e in hooks_bank.derived({})]
    assert "هوك المهمة الموجّهة" in " ".join(texts) and not any("حر" in t for t in texts)
    mine, other = stage3_generate.used_from_scripts("TESTP")
    assert all("هوك حر" not in h for h in other) and all("هوك حر" not in str(v) for v in mine.values())


def test_scripts_list_reports_the_mode_and_the_script_editor_refuses_free_projects(env):
    made = fp.create("p", rows("كلام"))
    env.script_file("g.json", make_script())
    from web import app as webapp
    c = TestClient(webapp.app)
    modes = {s["file"]: s["mode"] for s in c.get("/api/scripts").json()}
    assert modes[made["file"]] == "free" and modes["g.json"] == "guided"
    r = c.put(f"/api/scripts/{made['file']}", json={"title": "x"})
    assert r.status_code == 409 and "المشروع الحر" in r.json()["detail"]
    assert c.put("/api/scripts/g.json", json={"title": "x"}).status_code == 200


# ───────────────────────── المسارات ─────────────────────────
def client():
    from web import routes_free, routes_studio
    app = FastAPI()
    app.include_router(routes_free.router)
    app.include_router(routes_studio.router)
    return TestClient(app)


def xlsx_with(rows_):
    wb = openpyxl.load_workbook(io.BytesIO(scene_import.template()))
    for r in rows_:
        wb["المشاهد"].append(r)
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def test_create_route_blank_text_excel_and_conflicts(env):
    c = client()
    r = c.post("/api/free/projects", data={"title": "فاضي"})
    assert r.status_code == 200 and r.json()["file"].startswith("free__")
    r = c.post("/api/free/projects", data={"title": "نص", "text": "مشهد أول\n\nمشهد تاني"})
    f = r.json()["file"]
    assert r.status_code == 200 and len(c.get(f"/api/studio/{f}/free/board").json()["scenes"]) == 2
    r = c.post("/api/free/projects", data={"title": "اكسل"}, files={"file": ("s.xlsx", xlsx_with([["كلام", None, None, 4, "قريبة", None]]), "application/octet-stream")})
    assert r.status_code == 200 and c.get(f"/api/studio/{r.json()['file']}/free/board").json()["length_ms"] == 4000
    r = c.post("/api/free/projects", data={"text": "x"}, files={"file": ("s.xlsx", xlsx_with([["كلام"]]), "application/octet-stream")})
    assert r.status_code == 400 and "مش الاتنين" in r.json()["detail"]
    assert c.post("/api/free/projects", data={"text": "x" * 20_001}).status_code == 400
    assert c.post("/api/free/projects", files={"file": ("s.xlsx", b"junk", "application/octet-stream")}).status_code == 400
    assert c.post("/api/free/projects", data={"product_file": "nope.json"}).status_code == 404


def test_template_download_route(env):
    r = client().get("/api/scenes/template.xlsx")
    assert r.status_code == 200 and r.content[:2] == b"PK" and "spreadsheetml" in r.headers["content-type"]
    assert openpyxl.load_workbook(io.BytesIO(r.content)).sheetnames[:2] == ["اقرأني", "المشاهد"]


def test_structure_import_and_board_routes(env):
    c = client()
    f = c.post("/api/free/projects", data={"title": "p", "text": "أول\n\nتاني"}).json()["file"]
    b = c.get(f"/api/studio/{f}/free/board").json()
    assert b["readiness"]["no_media"] and b["limits"]["scenes"] == timeline.FREE_MAX_SCENES
    payload = {"rev": b["rev"], "scenes": [{"key": s["key"], "voiceover": "غيّرت", "shots": [{"id": s["shots"][0]["id"], "duration_ms": 4000}]} for s in b["scenes"]]}
    r = c.put(f"/api/studio/{f}/free/structure", json=payload)
    assert r.status_code == 200 and r.json()["length_ms"] == 8000
    assert c.put(f"/api/studio/{f}/free/structure", json=payload).status_code == 409                  # rev قديم
    assert c.put(f"/api/studio/{f}/free/structure", json={"rev": r.json()["rev"], "scenes": "x"}).status_code == 400
    r = c.post(f"/api/studio/{f}/free/import", data={"text": "تالت"})
    assert r.status_code == 200 and len(r.json()["board"]["scenes"]) == 3
    assert c.post(f"/api/studio/{f}/free/import", data={}).status_code == 400
    r = c.post(f"/api/studio/{f}/free/import", data={"text": "بس هاد", "replace": "true"})
    assert [s["voiceover"] for s in r.json()["board"]["scenes"]] == ["بس هاد"]
    assert c.get("/api/studio/none.json/free/board").status_code == 404


def test_media_routes_upload_images_to_new_scenes_and_to_a_shot(env, tmp_path):
    c = client()
    f = c.post("/api/free/projects", data={"title": "p"}).json()["file"]
    png = img_file(tmp_path).read_bytes()
    r = c.post(f"/api/studio/{f}/free/media", files=[("files", ("a.png", png, "image/png")), ("files", ("b.png", png, "image/png"))])
    assert r.status_code == 200 and len(r.json()["board"]["scenes"]) == 2 and r.json()["board"]["readiness"]["ready"]
    assert not (Path(config.LIBRARY_DIR) / "uploads" / "free_tmp").exists() or not any((Path(config.LIBRARY_DIR) / "uploads" / "free_tmp").iterdir())
    r = c.post(f"/api/studio/{f}/free/media", files=[("files", ("evil.exe", b"MZ", "application/octet-stream"))])
    assert r.status_code == 400 and "evil.exe" in r.json()["detail"]
    r = c.post(f"/api/studio/{f}/free/media", files=[("files", ("fake.png", b"not image", "image/png"))])
    assert r.status_code == 400 and "fake.png" in r.json()["detail"]
    b = c.get(f"/api/studio/{f}/free/board").json()
    sid = b["scenes"][0]["shots"][0]["id"]
    r = c.post(f"/api/studio/{f}/shots/{sid}/upload", files={"file": ("n.png", png, "image/png")})
    assert r.status_code == 200 and r.json()["shot"]["id"] == sid and len(r.json()["shot"]["image"]["candidates"]) == 2
    assert c.post(f"/api/studio/{f}/shots/{sid}/upload", files={"file": ("n.txt", b"x", "text/plain")}).status_code == 400
    assert c.post(f"/api/studio/{f}/shots/f999/upload", files={"file": ("n.png", png, "image/png")}).status_code == 404


def test_shot_upload_is_closed_for_guided_scripts_for_now(env, tmp_path):
    env.script_file("g.json", make_script())
    r = client().post("/api/studio/g.json/shots/s1c1/upload", files={"file": ("n.png", img_file(tmp_path).read_bytes(), "image/png")})
    assert r.status_code == 409


def test_image_generation_for_a_free_shot_needs_its_description_not_station_6(env):
    c = client()
    f = c.post("/api/free/projects", data={"title": "p", "text": "أول"}).json()["file"]
    sid = c.get(f"/api/studio/{f}/free/board").json()["scenes"][0]["shots"][0]["id"]
    r = c.post(f"/api/studio/{f}/images/generate", params={"shot": sid})
    assert r.status_code == 409 and "وصف الصورة" in r.json()["detail"]
    r = c.put(f"/api/studio/{f}/shots/{sid}/prompt", json={"prompt_ar": "علبة قطرة على طاولة خشب"})
    assert r.status_code == 200 and "بدون أي كتابة" in r.json()["shot"]["prompt_ar"]
    r = c.post(f"/api/studio/{f}/images/generate", params={"shot": sid})
    assert r.status_code == 200 and r.json()["job_id"].startswith("job")
    assert env.jobs[-1]["cmd"][:2] == ["image_studio.py", "gen"] and "--shot" in env.jobs[-1]["cmd"]
    assert c.post(f"/api/studio/{f}/images/generate", params={"shot": "f999"}).status_code == 404
    env.script_file("g.json", make_script())
    r = c.post("/api/studio/g.json/images/generate", params={"shot": "s1c1"})
    assert r.status_code == 409                                         # الموجّه: نفس شروطه (صوت معتمد…)


def test_other_studio_routes_work_on_a_free_timeline_without_voice_or_chain(env, tmp_path):
    """نفس مسارات اللقطات بالمهمة (برومبت/اختيار صورة) بتشتغل للحر بلا صوت معتمد ولا ترتيب محطات."""
    c = client()
    f = c.post("/api/free/projects", data={"title": "p", "text": "أول"}).json()["file"]
    sid = c.get(f"/api/studio/{f}/free/board").json()["scenes"][0]["shots"][0]["id"]
    assert c.get(f"/api/studio/{f}/state").json()["stations"][0]["state"] == "approved"
    r = c.post(f"/api/studio/{f}/shots/{sid}/upload", files={"file": ("a.png", img_file(tmp_path).read_bytes(), "image/png")})
    cand = r.json()["shot"]["image"]["candidates"][0]["path"]
    r = c.post(f"/api/studio/{f}/shots/{sid}/choose", json={"path": cand})
    assert r.status_code == 200 and r.json()["shot"]["image"]["chosen"].endswith("_1080.png")
    assert c.get(f"/api/studio/{f}/timeline").status_code == 200
