"""ورقة المنتج وورقة الشخصية: وضعيات ثابتة، لكل وضعية برومبت جاهز بنفس الإطار والنِّسب، صورة جنب برومبتها، تثبيت، فحص نسبة، Word."""

import io
import json
import os
import re
import sys
import zipfile
from pathlib import Path

import pytest
from PIL import Image

sys.path.insert(0, os.path.dirname(__file__))
from a_helpers import env  # noqa: E402,F401

import characters  # noqa: E402
import config  # noqa: E402
import product_ref  # noqa: E402
import refprompts  # noqa: E402

ROOT = Path(__file__).resolve().parent.parent
DIMS = {"h": 9, "w": 3.2, "d": 3.2}
LABEL = "ماء الكمأة قطرة عينية"


def cutout(w, h, canvas=(400, 500)):
    """PNG مفرَّغ: مستطيل معتم w×h (بالبكسل) وسط خلفية شفافة — إطار المنتج الفعلي = w×h بالضبط."""
    im = Image.new("RGBA", canvas, (0, 0, 0, 0))
    solid = Image.new("RGBA", (w, h), (30, 80, 200, 255))
    im.paste(solid, ((canvas[0] - w) // 2, (canvas[1] - h) // 2))
    b = io.BytesIO()
    im.save(b, "PNG")
    return b.getvalue()


def full_details(code="kama", **kw):
    return product_ref.set_details(code, kw.get("dims", DIMS), kw.get("label", LABEL), kw.get("colors", ["أبيض", "أزرق"]), kw.get("visual", "white 30 ml dropper bottle with a blue cap"))


def lines(text):
    return text.split("\n")


# ═════════════════════════ المنتج: كل منتج مستقل ═════════════════════════
def test_each_product_has_its_own_independent_reference_sheet(env):
    env.product("kama.json", product_name="KAMA", product_code="kama", category="eye_care", product_visual="عبوة قطرة بيضاء")
    env.product("vara.json", product_name="VARA", product_code="vara", category="hair_removal", product_visual="tube of cream")
    full_details("kama")
    product_ref.set_details("vara", {"h": 15, "w": 5, "d": 5}, "VARA CREAM", ["pink"])
    product_ref.set_angle("kama", "front", cutout(100, 280))
    product_ref.set_lock("kama", "front", True)
    product_ref.set_prompt("kama", "back", "my custom back prompt")
    rows = {r["code"]: r for r in product_ref.list_products()}
    assert set(rows) == {"kama", "vara"} and rows["kama"]["done"] == 1 and rows["vara"]["done"] == 0 and rows["kama"]["locked"] == 1 and rows["vara"]["locked"] == 0
    assert rows["kama"]["name"] == "KAMA" and rows["vara"]["required"] == 6 and rows["vara"]["complete"] is False
    k, v = product_ref.sheet("kama"), product_ref.sheet("vara")
    kf, vf = k["slots"][0]["prompt"]["text"], v["slots"][0]["prompt"]["text"]
    assert "KAMA" in kf and LABEL in kf and "9:3.2:3.2" in kf and "VARA" not in kf
    assert "VARA" in vf and "VARA CREAM" in vf and "15:5:5" in vf and LABEL not in vf and "KAMA" not in vf
    assert k["slots"][1]["prompt"]["edited"] is True and v["slots"][1]["prompt"]["edited"] is False       # التعديل بمنتج واحد بس
    assert (config.LIBRARY_DIR / "product_refs" / "kama" / "ref.json").exists() and (config.LIBRARY_DIR / "product_refs" / "vara" / "ref.json").exists()
    assert not (config.LIBRARY_DIR / "product_refs" / "vara" / "angles" / "front.png").exists()


def test_a_ref_folder_without_a_product_file_is_still_listed(env):
    product_ref.set_details("orphan", None, "X", None)
    assert [r["code"] for r in product_ref.list_products()] == ["orphan"]
    assert product_ref.list_products()[0]["name"] == "orphan"


# ═════════════════════════ البرومبتات: ثابتة إلا الوضعية ═════════════════════════
def test_prompts_share_frame_light_proportions_and_only_the_pose_changes(env):
    env.product("kama.json", product_name="KAMA", product_code="kama", category="eye_care")
    full_details()
    ref = product_ref.load("kama")
    ps = {s: product_ref.default_prompt(ref, s) for s in refprompts.PRODUCT_SLOT_KEYS}
    ang = ["front", "back", "left", "right", "top", "bottom"]
    framing = {next(l for l in lines(ps[s]) if l.startswith("Framing")) for s in ang}
    light = {next(l for l in lines(ps[s]) if l.startswith("Lighting")) for s in ang}
    props = {next(l for l in lines(ps[s]) if l.startswith("True proportions")) for s in ang}
    assert len(framing) == len(light) == len(props) == 1                                          # نفس الإطار والإضاءة والنِّسب بكل الزوايا
    assert "80% of the image height" in framing.pop() and "pure-white seamless" in light.pop() and "9:3.2:3.2" in props.pop()
    poses = [next(l for l in lines(ps[s]) if l.startswith("Camera/pose")) for s in ang]
    assert len(set(poses)) == 6                                                                  # وضعية الكاميرا هي الوحيدة المختلفة
    for s in ang:
        t = ps[s]
        assert f'"{LABEL}"' in t and "أبيض, أزرق" in t and "white 30 ml dropper bottle" in t and "H:W:D" in t
        assert ("IDENTITY LOCK" in t) == (s != "front")                                          # قفل الهوية بغير «أمام»
        assert "no hands" in t
    assert "FRONT-view image" in ps["back"] and "perfectly level" in ps["front"] and "directly above" in ps["top"] and "directly below" in ps["bottom"]
    assert "exactly 180°" in ps["back"] and "exactly 90°" in ps["left"] and "LEFT side" in ps["left"] and "RIGHT side" in ps["right"]
    assert "exactly one hand" in ps["in_hand"] and "no hands" not in ps["in_hand"] and "label area fills the frame" in ps["label"]


def test_missing_details_degrade_gracefully_and_are_listed_as_gaps(env):
    env.product("kama.json", product_name="KAMA", product_code="kama", category="eye_care", product_visual="")
    sh = product_ref.sheet("kama")
    t = sh["slots"][0]["prompt"]["text"]
    assert "True proportions" not in t and "Keep the product's proportions identical in every view" in t
    assert "Do not add any text to the product" in t and "Colors:" not in t and "Look (as defined" not in t
    assert sh["prompt_gaps"] == ["الأبعاد بالسنتيمتر (بتثبّت النِّسب والحجم)", "وصف شكل المنتج", "نص الملصق الحرفي", "الألوان"]
    full_details()
    assert product_ref.sheet("kama")["prompt_gaps"] == []


def test_visual_comes_from_the_ref_then_the_brief(env):
    env.product("kama.json", product_name="KAMA", product_code="kama", category="eye_care", product_visual="وصف البريف")
    sh = product_ref.sheet("kama")
    assert sh["visual_from_brief"] is True and sh["brief_visual"] == "وصف البريف" and "وصف البريف" in sh["slots"][0]["prompt"]["text"]
    assert sh["prompt_gaps"].count("وصف شكل المنتج") == 0                                          # وصف البريف بيكفي
    product_ref.set_details("kama", visual="own description")
    sh = product_ref.sheet("kama")
    assert sh["visual_from_brief"] is False and "own description" in sh["slots"][0]["prompt"]["text"] and "وصف البريف" not in sh["slots"][0]["prompt"]["text"]
    with pytest.raises(product_ref.ProductRefError, match="طويل"):
        product_ref.set_details("kama", visual="x" * 601)


def test_prompt_override_is_saved_flagged_and_resettable_and_defaults_follow_the_details(env):
    env.product("kama.json", product_name="KAMA", product_code="kama")
    full_details()
    d0 = product_ref.sheet("kama")["slots"][1]["prompt"]["default"]
    product_ref.set_prompt("kama", "back", "  BACK custom prompt  ")
    s = product_ref.sheet("kama")["slots"][1]["prompt"]
    assert s == {"text": "BACK custom prompt", "default": d0, "edited": True}
    assert json.loads((config.LIBRARY_DIR / "product_refs" / "kama" / "ref.json").read_text(encoding="utf-8"))["prompts"] == {"back": "BACK custom prompt"}
    product_ref.set_details("kama", {"h": 10, "w": 4, "d": 4})
    s = product_ref.sheet("kama")["slots"][1]["prompt"]
    assert s["text"] == "BACK custom prompt" and "10:4:4" in s["default"] and s["default"] != d0         # الأصلي بيتحدّث، وتعديلك بيضل
    product_ref.reset_prompt("kama", "back")
    s = product_ref.sheet("kama")["slots"][1]["prompt"]
    assert s["edited"] is False and s["text"] == s["default"]
    for bad, msg in (("", "فاضي"), ("   ", "فاضي"), ("x" * 4001, "أطول")):
        with pytest.raises(product_ref.ProductRefError, match=msg):
            product_ref.set_prompt("kama", "back", bad)
    for fn in (lambda: product_ref.set_prompt("kama", "nope", "x"), lambda: product_ref.reset_prompt("kama", "nope"), lambda: product_ref.set_lock("kama", "nope", True)):
        with pytest.raises(product_ref.ProductRefError, match="غير معروفة"):
            fn()


# ═════════════════════════ التثبيت ═════════════════════════
def test_lock_blocks_replacing_the_image_until_unlocked(env):
    full_details()
    product_ref.set_angle("kama", "front", cutout(100, 280))
    product_ref.set_lock("kama", "front", True)
    assert product_ref.sheet("kama")["slots"][0]["locked"] is True
    before = product_ref.angle_path("kama", "front").read_bytes()
    with pytest.raises(product_ref.ProductRefError, match="مثبّتة"):
        product_ref.set_angle("kama", "front", cutout(120, 300))
    assert product_ref.angle_path("kama", "front").read_bytes() == before                          # الصورة ما انلمست
    product_ref.set_angle("kama", "back", cutout(100, 280))                                         # باقي الوضعيات حرّة
    product_ref.set_lock("kama", "front", False)
    product_ref.set_angle("kama", "front", cutout(120, 300))
    assert product_ref.angle_path("kama", "front").read_bytes() != before
    with pytest.raises(product_ref.ProductRefError, match="ما في صورة"):
        product_ref.set_lock("kama", "left", True)


# ═════════════════════════ فحص النسبة بالأبعاد الحقيقية ═════════════════════════
def test_ratio_check_compares_each_view_with_the_real_dimensions(env):
    full_details(dims={"h": 9, "w": 3, "d": 6})
    product_ref.set_angle("kama", "front", cutout(100, 300))      # w/h = 0.333 ✓ (3/9)
    product_ref.set_angle("kama", "back", cutout(160, 300))       # 0.533 ✗
    product_ref.set_angle("kama", "left", cutout(200, 300))       # d/h = 0.667 ✓ (6/9)
    product_ref.set_angle("kama", "top", cutout(180, 300))        # w/d = 0.5 ✓ (3/6)? الصورة 0.6 ← فرق 20% ✗
    product_ref.set_angle("kama", "in_hand", cutout(150, 300))    # ما في توقّع
    ref = product_ref.load("kama")
    c = {s: product_ref.ratio_check(ref, s) for s in ("front", "back", "left", "right", "top", "in_hand")}
    assert c["front"]["ok"] is True and c["front"]["expected"] == 0.333 and c["front"]["actual"] == 0.333 and c["front"]["diff_pct"] == 0
    assert c["back"]["ok"] is False and c["back"]["diff_pct"] == 60
    assert c["left"]["ok"] is True and c["right"] is None and c["in_hand"] is None                   # يمين بلا صورة ← None، وبيد بلا توقّع
    assert c["top"]["ok"] is False and c["top"]["expected"] == 0.5 and c["top"]["actual"] == 0.6
    product_ref.set_details("kama", {"h": "", "w": "", "d": ""})
    nodims = product_ref.ratio_check(product_ref.load("kama"), "front")
    assert nodims["ok"] is None and nodims["expected"] is None and nodims["actual"] == 0.333         # الأبعاد ناقصة ← ما بنحكم
    assert product_ref.sheet("kama")["slots"][0]["check"]["ok"] is None


def test_tolerance_boundary_is_fifteen_percent(env):
    full_details(dims={"h": 10, "w": 5, "d": 5})
    product_ref.set_angle("kama", "front", cutout(57, 100))       # 0.57 مقابل 0.5 ← 14% ✓
    assert product_ref.ratio_check(product_ref.load("kama"), "front")["ok"] is True
    product_ref.set_angle("kama", "back", cutout(58, 100))        # 16% ✗
    assert product_ref.ratio_check(product_ref.load("kama"), "back")["ok"] is False


# ═════════════════════════ الوضعيات الاختيارية والتوافق ═════════════════════════
def test_optional_slots_do_not_affect_completeness_and_old_files_still_load(env):
    ref = product_ref.load("kama")
    assert list(ref["angles"]) == ["front", "back", "left", "right", "top", "bottom", "in_hand", "label", "box_front", "with_box"]
    assert [m for m in product_ref.missing(ref) if m.startswith("زاوية")] == [f"زاوية: {a}" for a in product_ref.ANGLES.values()]    # الإلزامية بس
    full_details()
    for a in product_ref.ANGLES:
        product_ref.set_angle("kama", a, cutout(100, 280))
    assert product_ref.missing(product_ref.load("kama")) == []                                       # بدون بيد/ملصق = كامل
    assert product_ref.list_products()[0]["complete"] is True and product_ref.list_products()[0]["extra_done"] == 0
    product_ref.set_angle("kama", "label", cutout(300, 200))
    assert product_ref.list_products()[0]["extra_done"] == 1
    # ملف قديم (6 زوايا فقط، بلا prompts/locked/meta/visual)
    legacy = {"code": "old", "angles": {a: None for a in product_ref.ANGLES}, "dims_cm": {"h": 1, "w": 1, "d": 1}, "label_text": "x", "colors": ["a"]}
    d = config.LIBRARY_DIR / "product_refs" / "old"
    d.mkdir(parents=True)
    (d / "ref.json").write_text(json.dumps(legacy), encoding="utf-8")
    ref = product_ref.load("old")
    assert "in_hand" in ref["angles"] and ref["prompts"] == {} and ref["locked"] == {} and ref["visual"] == ""
    assert len(product_ref.sheet("old")["slots"]) == 10


def test_invalid_codes_and_slots_are_rejected(env):
    with pytest.raises(product_ref.ProductRefError, match="كود"):
        product_ref.sheet("../")
    assert product_ref.load("../x")["code"] == "x"                                                 # الكود بيتنظّف (ما بيطلع من مجلد المنتجات)
    with pytest.raises(product_ref.ProductRefError, match="غير معروفة"):
        product_ref.set_angle("kama", "diagonal", cutout(10, 10))


# ═════════════════════════ الشخصية ═════════════════════════
def make_char(**kw):
    d = {"name_ar": "ليلى", "gender": "f", "age": 32, "skin": "wheat tone", "hair": "long black wavy", "persona": "warm and trustworthy",
         "wardrobe": {"description": "beige silk blouse", "palette": ["beige", "cream"]}, "market_look": {"SY": "Levantine features"},
         "categories": ["eye_care"], "voice": {"provider": "elevenlabs", "voice_id": "v1", "model": "m"}}
    d.update(kw)
    return characters.create(d)


def test_character_prompts_keep_one_identity_across_views_and_only_the_pose_changes(env):
    ch = make_char()
    sh = characters.slots(ch["id"], "SY")
    assert [s["slot"] for s in sh["slots"]] == ["front", "side", "three_quarter", "full_body", "hands"] and all(s["required"] for s in sh["slots"])
    texts = {s["slot"]: s["prompt"]["text"] for s in sh["slots"]}
    for v, t in texts.items():
        assert "32-year-old woman, warm and trustworthy" in t and "Skin: wheat tone" in t and "Hair: long black wavy" in t
        assert "Regional look: Levantine features" in t and "beige silk blouse" in t and "Palette: beige, cream" in t
        assert ("IDENTITY LOCK" in t) == (v != "front")
        assert next(l for l in lines(t) if l.startswith("Background & light")) == next(l for l in lines(texts["front"]) if l.startswith("Background & light"))
    assert len({next(l for l in lines(t) if l.startswith("Pose:")) for t in texts.values()}) == 5
    assert "pure side profile" in texts["side"] and "45°" in texts["three_quarter"] and "feet visible" in texts["full_body"] and "both hands" in texts["hands"]
    assert sh["prompt_gaps"] == [] and sh["name"] == "ليلى" and sh["countries"] == ["SY"]
    assert {s["slot"]: s["use_note"] for s in sh["slots"]} == {"front": "بتنبعت مع لقطات الوجه", "side": "", "three_quarter": "بتنبعت مع لقطات الوجه",
                                                                "full_body": "", "hands": "بتنبعت مع لقطات اليد بالمنتج"}
    male = make_char(name_ar="سامر", gender="m", age=40)
    assert "40-year-old man" in characters.slots(male["id"])["slots"][0]["prompt"]["text"]


def test_character_market_look_picks_the_requested_country_or_the_only_one(env):
    ch = make_char(market_look={"SY": "shami", "SA": "gulf"})
    sy, sa, none = (characters.slots(ch["id"], c)["slots"][0]["prompt"]["text"] for c in ("SY", "SA", ""))
    assert "Regional look: shami" in sy and "gulf" not in sy and "Regional look: gulf" in sa
    assert "SY: shami" in none and "SA: gulf" in none                                                 # بدون بلد: لحد 2 مع الرموز
    one = make_char(market_look={"SY": "only look"})
    assert "Regional look: only look" in characters.slots(one["id"])["slots"][0]["prompt"]["text"]


def test_real_person_prompts_use_the_attached_photos_as_the_only_identity_source(env):
    ch = make_char(source="real_person")
    t = characters.slots(ch["id"])["slots"][1]["prompt"]["text"]
    assert "IDENTITY SOURCE" in t and "ONLY source of identity" in t
    assert "IDENTITY SOURCE" not in characters.slots(make_char()["id"])["slots"][0]["prompt"]["text"]


def test_character_gaps_override_reset_and_lock(env):
    ch = make_char(skin="", hair="", persona="", wardrobe={"description": "", "palette": []})
    cid = ch["id"]
    assert characters.slots(cid)["prompt_gaps"] == ["البشرة", "الشعر", "الطابع", "وصف اللبس"]
    d0 = characters.slots(cid)["slots"][0]["prompt"]["default"]
    characters.set_prompt(cid, "front", "custom front")
    assert characters.slots(cid)["slots"][0]["prompt"] == {"text": "custom front", "default": d0, "edited": True}
    with pytest.raises(characters.CharacterError, match="فاضي"):
        characters.set_prompt(cid, "front", " ")
    with pytest.raises(characters.CharacterError, match="أطول"):
        characters.set_prompt(cid, "front", "x" * 4001)
    characters.update(cid, {"skin": "fair"})                                                          # تعديل بيانات ما بيمحي البرومبت المعدَّل
    assert characters.slots(cid)["slots"][0]["prompt"]["text"] == "custom front" and "Skin: fair" in characters.slots(cid)["slots"][1]["prompt"]["text"]
    characters.reset_prompt(cid, "front")
    assert characters.slots(cid)["slots"][0]["prompt"]["edited"] is False
    img = b"\x89PNG" + b"0" * 600
    with pytest.raises(characters.CharacterError, match="ما في صورة"):
        characters.set_lock(cid, "front", True)
    characters.set_sheet(cid, "front", "f.png", img)
    characters.set_lock(cid, "front", True)
    assert characters.slots(cid)["slots"][0]["locked"] is True
    with pytest.raises(characters.CharacterError, match="مثبّتة"):
        characters.set_sheet(cid, "front", "f2.png", img + b"1")
    assert characters.sheet_path(cid, "front").read_bytes() == img
    characters.set_sheet(cid, "side", "s.png", img)                                                    # غيرها حرّة
    characters.set_lock(cid, "front", False)
    characters.set_sheet(cid, "front", "f2.jpg", img + b"1")
    for fn in (lambda: characters.set_prompt(cid, "x", "t"), lambda: characters.reset_prompt(cid, "x"), lambda: characters.set_lock(cid, "x", True)):
        with pytest.raises(characters.CharacterError, match="غير معروفة"):
            fn()


def test_characters_without_the_new_keys_still_work(env):
    ch = make_char()
    p = characters.dir_of(ch["id"]) / "character.json"
    d = json.loads(p.read_text(encoding="utf-8"))
    d.pop("prompts"), d.pop("locked")
    p.write_text(json.dumps(d), encoding="utf-8")
    assert len(characters.slots(ch["id"])["slots"]) == 5
    characters.set_prompt(ch["id"], "side", "ok")
    characters.set_lock(ch["id"], "front", False)


# ═════════════════════════ Word ═════════════════════════
def read_docx(data: bytes):
    from docx import Document
    doc = Document(io.BytesIO(data))
    return doc, "\n".join(p.text for p in doc.paragraphs)


def test_product_prompts_docx_has_every_slot_prompt_and_arabic_rtl_headings(env):
    env.product("kama.json", product_name="KAMA", product_code="kama", category="eye_care")
    full_details()
    product_ref.set_prompt("kama", "left", "EDITED LEFT PROMPT")
    name, data = product_ref.prompts_docx("kama")
    assert name == "kama_reference_prompts.docx" and data[:2] == b"PK" and "word/document.xml" in zipfile.ZipFile(io.BytesIO(data)).namelist()
    doc, text = read_docx(data)
    assert "برومبتات الصور المرجعية — KAMA" in text and "ترتيب الشغل" in text and "أمام — FRONT" in text and "(اختيارية)" in text
    for s in refprompts.PRODUCT_SLOTS:
        if s["slot"] != "left":
            assert product_ref.default_prompt(product_ref.load("kama"), s["slot"]).split("\n")[0] in text
    assert "EDITED LEFT PROMPT" in text and "برومبت معدَّل" in text and "ارفق صورة «أمام» معه" in text
    ar = [p for p in doc.paragraphs if p.text.startswith("برومبتات الصور")]
    assert ar and ar[0]._p.xml.count("w:bidi") >= 1
    en = next(p for p in doc.paragraphs if p.text.startswith("PRODUCT REFERENCE SHEET — FRONT"))
    assert "w:bidi" not in en._p.xml and en.runs[0].font.name == "Consolas"                              # البرومبت LTR بخط أحادي


def test_docx_lists_the_missing_details_and_the_character_variant(env):
    env.product("kama.json", product_name="KAMA", product_code="kama")
    _, text = read_docx(product_ref.prompts_docx("kama")[1])
    assert "ناقص بالبرومبتات" in text and "الأبعاد بالسنتيمتر" in text
    ch = make_char(source="real_person")
    name, data = characters.prompts_docx(ch["id"], "SY")
    _, text = read_docx(data)
    assert name == f"{ch['id']}_reference_prompts.docx" and "برومبتات الورقة المرجعية — ليلى" in text and "مبنية على شخص حقيقي" in text
    assert text.count("CHARACTER REFERENCE SHEET") == 5 and "Regional look: Levantine features" in text


# ═════════════════════════ المسارات ═════════════════════════
@pytest.fixture
def client(env):
    from fastapi import FastAPI
    from fastapi.testclient import TestClient
    from web import routes_characters
    app = FastAPI()
    app.include_router(routes_characters.router)
    env.product("kama.json", product_name="KAMA", product_code="kama", category="eye_care")
    return TestClient(app)


def test_product_routes_list_sheet_prompt_lock_and_docx(client):
    assert [(p["code"], p["file"]) for p in client.get("/api/product-refs").json()["products"]] == [("kama", "kama.json")]
    r = client.put("/api/product-refs/kama", json={"dims_cm": DIMS, "label_text": LABEL, "colors": ["أبيض"], "visual": "white dropper bottle"})
    assert r.status_code == 200 and r.json()["visual"] == "white dropper bottle"
    d = client.get("/api/product-refs/kama").json()
    assert [s["slot"] for s in d["slots"]] == refprompts.PRODUCT_SLOT_KEYS and d["slot_labels"]["in_hand"] and d["angle_labels"]["front"] == "أمام" and "angles" in d
    assert d["product"]["name"] == "KAMA" and "white dropper bottle" in d["slots"][0]["prompt"]["text"] and d["prompt_gaps"] == []
    up = client.post("/api/product-refs/kama/angle/front", files={"file": ("f.png", cutout(100, 280), "image/png")})
    assert up.status_code == 200 and up.json()["angles"]["front"]
    assert client.post("/api/product-refs/kama/lock/front", json={"locked": True}).json()["slots"][0]["locked"] is True
    blocked = client.post("/api/product-refs/kama/angle/front", files={"file": ("f.png", cutout(100, 280), "image/png")})
    assert blocked.status_code == 400 and "مثبّتة" in blocked.json()["detail"]
    assert client.post("/api/product-refs/kama/lock/left", json={"locked": True}).status_code == 400
    s = client.put("/api/product-refs/kama/prompt/back", json={"text": "BACK mine"}).json()["slots"][1]["prompt"]
    assert s["text"] == "BACK mine" and s["edited"] is True
    assert client.put("/api/product-refs/kama/prompt/back", json={"text": ""}).status_code == 400
    assert client.put("/api/product-refs/kama/prompt/zzz", json={"text": "x"}).status_code == 400
    s = client.delete("/api/product-refs/kama/prompt/back").json()["slots"][1]["prompt"]
    assert s["edited"] is False
    doc = client.get("/api/product-refs/kama/prompts.docx")
    assert doc.status_code == 200 and doc.content[:2] == b"PK" and doc.headers["content-type"] == refprompts.DOCX_MIME
    assert 'filename="kama_reference_prompts.docx"' in doc.headers["content-disposition"]
    assert client.get("/api/product-refs/!!!/prompts.docx").status_code == 400
    assert client.get("/api/product-refs/kama/angle/in_hand").status_code == 404                           # الاختيارية بلا صورة


def test_character_routes_slots_prompt_lock_and_docx(client):
    cid = client.post("/api/characters", json={"name_ar": "ليلى", "gender": "f", "age": 30, "skin": "s", "hair": "h", "persona": "p",
                                               "wardrobe": {"description": "w", "palette": []}}).json()["id"]
    s = client.get(f"/api/characters/{cid}/slots").json()
    assert [x["slot"] for x in s["slots"]] == refprompts.CHARACTER_SLOT_KEYS and s["name"] == "ليلى" and s["prompt_gaps"] == []
    up = client.post(f"/api/characters/{cid}/sheet/front", files={"file": ("f.png", b"\x89PNG" + b"0" * 600, "image/png")})
    assert up.status_code == 200
    assert client.post(f"/api/characters/{cid}/lock/front", json={"locked": True}).json()["slots"][0]["locked"] is True
    again = client.post(f"/api/characters/{cid}/sheet/front", files={"file": ("g.png", b"\x89PNG" + b"1" * 600, "image/png")})
    assert again.status_code == 400 and "مثبّتة" in again.json()["detail"]
    assert client.put(f"/api/characters/{cid}/prompt/side", json={"text": "my side prompt"}).json()["slots"][1]["prompt"]["edited"] is True
    assert client.put(f"/api/characters/{cid}/prompt/side", json={"text": "x" * 4001}).status_code == 400
    assert client.delete(f"/api/characters/{cid}/prompt/side").json()["slots"][1]["prompt"]["edited"] is False
    assert client.get("/api/characters/ch_f_999/slots").status_code == 400
    doc = client.get(f"/api/characters/{cid}/prompts.docx?country=SY")
    assert doc.status_code == 200 and doc.content[:2] == b"PK" and f'filename="{cid}_reference_prompts.docx"' in doc.headers["content-disposition"]


# ═════════════════════════ الواجهة ═════════════════════════
def src(name):
    return (ROOT / "web" / "static" / name).read_text(encoding="utf-8")


def test_ui_shared_slot_component_is_a_closed_scope_and_escapes_everything():
    import test_ui_globals
    js = src("refslots.js")
    assert test_ui_globals.is_closed_scope(js) and "window.RefSlots = { render }" in js
    for must in ("${esc(s.label)}", "${esc(s.prompt.text)}", "${esc(cfg.imgUrl(s.slot))}", "${esc(s.pose_ar || \"\")}", "${esc(c.actual)}", "${esc(c.expected)}", "${esc(c.diff_pct)}", "${esc(s.slot)}"):
        assert must in js, must
    assert "confirm(" in js and "navigator.clipboard.writeText" in js and "readonly" in js
    for ev in ("onUpload", "onLock", "onSavePrompt", "onResetPrompt"):
        assert f"cfg.{ev}(" in js
    assert not re.search(r"[٠-٩]", js)
    for b in ("نسخ البرومبت", "عدّل", "احفظ البرومبت", "رجّع الأصلي", "ثبّت", "فكّ التثبيت", "مثبّتة", "اختيارية", "ارفق صورة «أمام»",
              "نزّل الصورة", "بتنبعت مع التوليد", "${esc(s.use_note)}", "cfg.fileBase"):
        assert b in js, b


def test_ui_product_and_character_screens_use_the_slots_and_every_server_text_is_escaped():
    pr, ch, idx = src("productref.js"), src("characters.js"), src("index.html")
    assert "RefSlots.render(" in pr and "RefSlots.render(" in ch
    assert "/api/product-refs" in pr and "prompts.docx" in pr and "prompts.docx" in ch and "${esc(d.product.name)}" in pr
    for must in ("/lock/", "/prompt/", "?auto=1", "d.prompt_gaps", "brief_visual", 'data-pr="visual"', "الوضعيات الثابتة"):
        assert must in pr, must
    assert "/slots" in ch and "chDrawSlots" in ch and "/lock/" in ch and "/prompt/" in ch and 'type="file"][data-view]' not in ch   # أزرار الرفع القديمة انشالت
    assert "CH_COUNTRY" in ch and 'id="chCountry"' in ch and "?country=" in ch and "/api/ref-sheets/prompts.docx" in ch          # ملامح بلد وحدة بالبرومبت
    assert idx.index("/refslots.js") < idx.index("/characters.js") < idx.index("/productref.js")
    assert not re.search(r"[٠-٩]", pr + ch)
    for name in ("slotrow", "slotlist", "prgrid", "prcard", "is-locked"):
        assert f".{name}" in src("style.css"), name
    assert "direction: ltr" in src("style.css")


def test_the_sheet_renders_inside_any_container_without_global_ids_and_the_old_refs_tab_is_gone():
    pr, app, idx = src("productref.js"), src("app.js"), src("index.html")
    body = pr[pr.index("async function prMountSheet"):]
    assert '$("#' not in body and "getElementById" not in body                                                       # الورقة بتنعرض بمكانين: ما في id عام يتضارب
    assert "holder._prSeq" in pr and "prMountProductTab" in pr and "x.file === file" in pr                           # آخر نداء بيرسم + ربط المنتج بملفه
    assert "async function loadRefsPage() { await prMountProductTab(); }" in app
    assert 'id="prodRefSheet"' in idx and 'id="prod-refs"' in idx
    for gone in ("refGallery", "refUpload", "refGenBtn", "refGenRun", "refProduct", "REF_KIND_AR"):
        assert gone not in idx and gone not in app, gone                                                                # المعرض المخلوط والرفع القديم انشالوا
    assert "بتنبعت مع كل طلب توليد" not in idx                                                                          # النص الغلط (ما كان صحيح)
    assert "/api/assets" not in pr and "from-asset" not in pr and "prDrawLegacy" not in pr                         # الصور القديمة انشالت من الواجهة
    assert "onDelete" in pr and 'method: "DELETE"' in pr and 'data-pr="box"' in pr and "box_text" in pr           # حذف الصور + نص الكرتونة
    assert "libsChTabs" not in idx and "libsch-" not in idx and "libsch-" not in src("shell.js") and 'id="v-product-ref"' not in idx
    assert "prFileForCode" in pr and "prFileForCode" in src("shell.js") and "onDelete" in src("characters.js") and "data-del" in src("refslots.js")


def test_ui_mock_files_have_the_new_shapes():
    d = json.loads(src("mock/product_ref.json"))
    assert [s["slot"] for s in d["slots"]] == refprompts.PRODUCT_SLOT_KEYS and d["slots"][0]["prompt"]["text"].startswith("PRODUCT REFERENCE SHEET")
    assert any(s["locked"] for s in d["slots"]) and any(s["prompt"]["edited"] for s in d["slots"]) and any(s["check"] and s["check"]["ok"] is False for s in d["slots"])
    assert {s["slot"] for s in d["slots"] if s["use_note"]} == {"front"}
    lst = json.loads(src("mock/product_refs.json"))
    assert {p["code"] for p in lst["products"]} == {"kama", "vara"} and all("done" in p for p in lst["products"])
    assert {p["file"] for p in lst["products"]} == {"kama.json", "vara.json"}
    c = json.loads(src("mock/char_slots.json"))
    assert [s["slot"] for s in c["slots"]] == refprompts.CHARACTER_SLOT_KEYS
    app = src("app.js")
    assert 'path === "/api/product-refs") return mockFile("product_refs")' in app and 'mockFile("char_slots")' in app


def test_one_word_file_for_all_products_and_characters_numbering_restarts_per_section(env, client):
    import ref_docx
    env.product("vara.json", product_name="VARA", product_code="vara", category="hair_removal")
    full_details("kama")
    product_ref.set_prompt("kama", "left", "EDITED LEFT")
    product_ref.set_angle("kama", "front", cutout(100, 280))
    product_ref.set_lock("kama", "front", True)
    make_char()
    name, data = ref_docx.all_prompts_docx()
    assert name == "all_reference_prompts.docx"
    doc, text = read_docx(data)
    assert "برومبتات الصور المرجعية — كل المنتجات والشخصيات" in text and "المنتج: KAMA" in text and "المنتج: VARA" in text and "الشخصية: ليلى" in text
    assert text.index("المنتج: KAMA") < text.index("المنتج: VARA") < text.index("الشخصية: ليلى")                   # المنتجات أول وبعدها الشخصيات
    assert text.count("PRODUCT REFERENCE SHEET —") == 19 and text.count("CHARACTER REFERENCE SHEET") == 5        # 20 وضعية منتج بس واحدة منها معدَّلة بنص بلا ترويسة
    assert "EDITED LEFT" in text and "ناقص بالبرومبتات" in text and "المعرّف: ch_f_001" in text
    nums = [p.text.split(".")[0] for p in doc.paragraphs if re.match(r"^\d+\. ", p.text)]
    assert nums == [str(i) for i in range(1, 11)] * 2 + [str(i) for i in range(1, 6)]                           # الترقيم من 1 بكل قسم: 10 + 10 + 5
    assert "الصورة: مثبّتة" in text and "الصورة: ناقصة" in text and "بتنبعت مع كل لقطة" in text and "بتنبعت مع لقطات الوجه" in text
    r = client.get("/api/ref-sheets/prompts.docx")
    assert r.status_code == 200 and r.content[:2] == b"PK" and 'filename="all_reference_prompts.docx"' in r.headers["content-disposition"]
    assert "/api/ref-sheets/prompts.docx" in src("productref.js")


def test_all_prompts_docx_with_nothing_yet_still_opens_and_says_so(env):
    import ref_docx
    _, text = read_docx(ref_docx.all_prompts_docx()[1])
    assert "ما في منتجات ولا شخصيات بعد" in text


# ═════════════════════════ كود المنتج: طريقة تنظيف وحدة للورقة وللتوليد ═════════════════════════
def test_product_code_is_normalized_one_way_for_the_sheet_and_for_generation(env):
    import image_studio
    env.product("hair_removal.json", product_name="Hair Removal", product_code="", category="hair_removal")
    rows = {r["code"]: r for r in product_ref.list_products()}
    assert "hairremoval" in rows and rows["hairremoval"]["file"] == "hair_removal.json" and "hair_removal" not in rows
    gen_code = image_studio.product_code({"_meta": {"product": "Hair Removal", "product_file": "hair_removal.json"}})
    assert gen_code == "hairremoval" == product_ref._code("hair_removal") == product_ref._code("Hair-Removal")      # التوليد والورقة بيلاقوا نفس المجلد
    assert product_ref.dir_of("hair_removal") == product_ref.dir_of(gen_code)
    assert product_ref.product_info("hair_removal")["name"] == "Hair Removal"                                    # الكود القديم بيضل بيلاقي المنتج
    product_ref.set_angle(gen_code, "front", cutout(100, 280))
    assert [p.name for p in image_studio.refs_for_shot("s1", {"type": "close"}, gen_code)] == ["front.png"]
    assert len(product_ref._norm("x" * 40)) == 24                                                                   # نفس حد الاستوديو (24)


# ═════════════════════════ شو بيوصل للتوليد فعلياً (الشارة مربوطة بالمنطق الحقيقي) ═════════════════════════
def test_use_notes_match_what_the_image_studio_really_attaches(env):
    import image_studio
    env.product("kama.json", product_name="KAMA", product_code="kama", category="eye_care")
    ch = make_char()
    for v in characters.VIEWS:
        characters.set_sheet(ch["id"], v, f"{v}.png", b"\x89PNG" + b"0" * 600)
    for a in product_ref.SLOTS:
        product_ref.set_angle("kama", a, cutout(100, 280))
    sent_char, sent_prod = set(), set()
    for typ in ("close", "wide", "medium", "hand_product"):
        for p, role in image_studio.refs_with_roles("s1", {"type": typ, "refs": {"character_id": ch["id"]}}, "kama"):
            (sent_char if role == "person" else sent_prod).add(p.stem)
    noted_c = {s["slot"] for s in characters.slots(ch["id"])["slots"] if s["use_note"]}
    noted_p = {s["slot"] for s in product_ref.sheet("kama")["slots"] if s["use_note"]}
    assert sent_char == noted_c == {"front", "three_quarter", "hands"} and sent_prod == noted_p == {"front"}


def test_image_rule_describes_each_reference_by_its_role(tmp_path, monkeypatch):
    import imagegen
    assert imagegen.ref_rule(None) == imagegen.REF_RULE and imagegen.ref_rule([]) == imagegen.REF_RULE
    only_product = imagegen.ref_rule(["product"])
    assert "the person" not in only_product and "Reference image(s) 1: the product" in only_product         # سيناريو بلا شخصية: ما بنقول «أول صورة شخص»
    both = imagegen.ref_rule(["person", "person", "product"])
    assert "Reference image(s) 1, 2: the person" in both and "Reference image(s) 3: the product" in both
    seen = {}

    def fake(cfg, prompt, out, refs):
        seen["prompt"], seen["refs"] = prompt, [r.name for r in refs]
        out.write_bytes(b"x" * 2000)

    monkeypatch.setattr(imagegen, "image_config", lambda ov=None: {"provider": "p", "kind": "fake", "model": "m", "api_key": "k", "style": "STYLE", "supports_refs": True})
    monkeypatch.setitem(imagegen._DISPATCH, "fake", fake)
    a, b = tmp_path / "a.png", tmp_path / "b.png"
    a.write_bytes(b"1")
    b.write_bytes(b"2")
    imagegen.generate("shot", tmp_path / "o.png", refs=[a, tmp_path / "gone.png", b], ref_roles=["person", "person", "product"])
    assert seen["refs"] == ["a.png", "b.png"]                                                                   # صورة مفقودة ما بتزحزح أدوار الباقي
    assert "Reference image(s) 1: the person" in seen["prompt"] and "Reference image(s) 2: the product" in seen["prompt"]
    imagegen.generate("shot", tmp_path / "o2.png", refs=[a])
    assert imagegen.REF_RULE in seen["prompt"]                                                                  # بلا أدوار: الصيغة القديمة
    imagegen.generate("shot", tmp_path / "o3.png")
    assert imagegen.REF_RULE not in seen["prompt"] and "Reference image" not in seen["prompt"]


# ═════════════════════════ الكرتونة ═════════════════════════
def test_box_slots_have_their_own_text_and_prompts_and_never_block_completeness(env):
    env.product("kama.json", product_name="KAMA", product_code="kama", category="eye_care")
    full_details()
    product_ref.set_details("kama", box_text="كاما — قطرة عينية 30 مل")
    ref = product_ref.load("kama")
    assert ref["box_text"] == "كاما — قطرة عينية 30 مل"
    bf, wb = product_ref.default_prompt(ref, "box_front"), product_ref.default_prompt(ref, "with_box")
    assert "Box text — reproduce it letter for letter" in bf and "كاما — قطرة عينية 30 مل" in bf and "Label text" not in bf and "True proportions" not in bf
    assert "retail carton" in bf and "one retail box only" in bf and "IDENTITY LOCK" in bf
    assert "Label text" in wb and "Box text" in wb and "True proportions: height 9 cm" in wb and "one product and its box only" in wb
    product_ref.set_details("kama", box_text="")
    assert "Do not add any text to the box" in product_ref.default_prompt(product_ref.load("kama"), "box_front")
    with pytest.raises(product_ref.ProductRefError, match="طويل"):
        product_ref.set_details("kama", box_text="x" * 601)
    for a in product_ref.ANGLES:                                                                      # اختيارية: ما بتأثر على «كامل» وما فيها فحص نسبة
        product_ref.set_angle("kama", a, cutout(100, 280))
    product_ref.set_angle("kama", "box_front", cutout(200, 120))
    row = product_ref.list_products()[0]
    assert row["complete"] is True and row["extra_done"] == 1
    slot = product_ref.sheet("kama")["slots"][8]
    assert slot["slot"] == "box_front" and slot["check"] is None and slot["required"] is False


# ═════════════════════════ حذف الصور (للسلة، مش مسح نهائي) ═════════════════════════
def test_product_image_delete_moves_it_to_the_trash_and_respects_the_lock(env):
    full_details()
    product_ref.set_angle("kama", "front", cutout(100, 280))
    product_ref.set_angle("kama", "back", cutout(100, 280))
    product_ref.set_lock("kama", "front", True)
    with pytest.raises(product_ref.ProductRefError, match="مثبّتة"):
        product_ref.delete_angle("kama", "front")
    assert product_ref.angle_path("kama", "front")
    ref = product_ref.delete_angle("kama", "back")
    assert ref["angles"]["back"] is None and "back" not in ref["meta"] and product_ref.angle_path("kama", "back") is None
    trash = list((config.LIBRARY_DIR / "_trash" / "product_refs" / "kama").glob("back_*.png"))
    assert len(trash) == 1 and trash[0].stat().st_size > 0                                            # بالسلة مش ممسوحة
    with pytest.raises(product_ref.ProductRefError, match="ما في صورة"):
        product_ref.delete_angle("kama", "back")
    with pytest.raises(product_ref.ProductRefError, match="غير معروفة"):
        product_ref.delete_angle("kama", "nope")
    assert "زاوية: خلف" in product_ref.missing(product_ref.load("kama"))


def test_character_image_delete_goes_to_the_trash_and_drops_a_ready_character_to_draft(env):
    ch = make_char()
    img = b"\x89PNG" + b"0" * 600
    for v in characters.VIEWS:
        characters.set_sheet(ch["id"], v, f"{v}.png", img)
    assert characters.get(ch["id"])["status"] == "ready"
    characters.set_lock(ch["id"], "side", True)
    with pytest.raises(characters.CharacterError, match="مثبّتة"):
        characters.delete_sheet(ch["id"], "side")
    out = characters.delete_sheet(ch["id"], "hands")
    assert out["sheet"]["hands"] is None and out["status"] == "draft"
    assert characters.sheet_path(ch["id"], "hands") is None
    assert len(list((config.LIBRARY_DIR / "_trash" / "characters" / ch["id"]).glob("hands_*.png"))) == 1
    assert "ورقة مرجعية: يدين" in characters.missing_for_ready(characters.get(ch["id"]))
    with pytest.raises(characters.CharacterError, match="ما في صورة"):
        characters.delete_sheet(ch["id"], "hands")


def test_delete_routes_for_product_and_character_images(client):
    png = ("f.png", cutout(100, 280), "image/png")
    assert client.post("/api/product-refs/kama/angle/front", files={"file": png}).status_code == 200
    r = client.delete("/api/product-refs/kama/angle/front")
    assert r.status_code == 200 and r.json()["angles"]["front"] is None and r.json()["slots"][0]["has_image"] is False
    assert client.delete("/api/product-refs/kama/angle/front").status_code == 400                      # ما في صورة
    client.post("/api/product-refs/kama/angle/back", files={"file": png})
    client.post("/api/product-refs/kama/lock/back", json={"locked": True})
    blocked = client.delete("/api/product-refs/kama/angle/back")
    assert blocked.status_code == 400 and "مثبّتة" in blocked.json()["detail"]
    cid = client.post("/api/characters", json={"name_ar": "ليلى", "gender": "f", "age": 30, "skin": "s", "hair": "h", "persona": "p",
                                               "wardrobe": {"description": "w", "palette": []}}).json()["id"]
    assert client.post(f"/api/characters/{cid}/sheet/front", files={"file": ("f.png", b"\x89PNG" + b"0" * 600, "image/png")}).status_code == 200
    d = client.delete(f"/api/characters/{cid}/sheet/front")
    assert d.status_code == 200 and d.json()["sheet"]["front"] is None
    assert client.delete(f"/api/characters/{cid}/sheet/front").status_code == 400
    assert client.delete("/api/characters/ch_f_999/sheet/front").status_code == 400
