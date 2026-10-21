"""استيراد الأنماط من Excel: النموذج، المعاينة (صح/غلط/مشابه)، التنفيذ، وحماية الملف."""

import io
import json
import os
import sys
import time
import zipfile

import openpyxl
import pytest

sys.path.insert(0, os.path.dirname(os.path.dirname(__file__)))

import clusters  # noqa: E402
import config  # noqa: E402
import pattern_excel as px  # noqa: E402
import pattern_fields as pf  # noqa: E402
import pattern_store as ps  # noqa: E402
import review_queue  # noqa: E402


@pytest.fixture
def lib(tmp_path, monkeypatch):
    pdir = tmp_path / "library" / "patterns"
    pdir.mkdir(parents=True)
    monkeypatch.setattr(config, "PATTERNS_DIR", pdir)
    monkeypatch.setattr(config, "SCRIPTS_DIR", tmp_path / "library" / "scripts")
    monkeypatch.setattr(review_queue, "QUEUE_PATH", tmp_path / "rq.json")
    monkeypatch.setattr(clusters, "CLUSTERS_PATH", tmp_path / "library" / "clusters.json")
    return tmp_path


def label_to_col(ws):
    return {str(c.value).replace(" *", "").strip(): c.column for c in ws[1]}


def make_xlsx(patterns=None, scenes=None, *, tweak=None) -> bytes:
    """نموذج معبّى: patterns = [{path أو «عنوان»: قيمة}]، scenes = [{«اسم النمط»…}]."""
    wb = openpyxl.load_workbook(io.BytesIO(px.template()))
    wp, ws = wb[px.PATTERNS_SHEET], wb[px.SCENES_SHEET]
    cols = {spec["path"]: i for i, spec in enumerate(pf.FIELDS, 1)}
    for r, row in enumerate(patterns or [], 2):
        for path, v in row.items():
            wp.cell(row=r, column=cols[path], value="\n".join(v) if isinstance(v, list) else v)
    scols = label_to_col(ws)
    for r, row in enumerate(scenes or [], 2):
        for k, v in row.items():
            ws.cell(row=r, column=scols[k], value=v)
    if tweak:
        tweak(wb)
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def example_pattern(**over):
    fields, scenes = px._example()
    row = {**fields, **over}
    name = row["pattern_name"]
    sc = [{"اسم النمط": name, "رقم المشهد": i, "وظيفة المشهد": s["function"], "المدة (ثانية)": s["duration_sec"], "نوع اللقطة": s["shot_type"],
           "الهدف من الصوت": s["vo_intent"], "دور النص عالشاشة": s["on_screen_text_role"], "الإحساس": s["emotion"]} for i, s in enumerate(scenes, 1)]
    return row, sc


# ═════════════════ النموذج ═════════════════
def test_template_has_the_three_sheets_a_hidden_list_sheet_and_dropdowns():
    wb = openpyxl.load_workbook(io.BytesIO(px.template()))
    assert wb.sheetnames == [px.README_SHEET, px.PATTERNS_SHEET, px.SCENES_SHEET, px.LISTS_SHEET]
    assert wb[px.LISTS_SHEET].sheet_state == "hidden"
    heads = [c.value for c in wb[px.PATTERNS_SHEET][1]]
    assert heads[0] == "اسم النمط *" and len(heads) == len(pf.FIELDS)
    assert [c.value for c in wb[px.SCENES_SHEET][1]][:2] == ["اسم النمط *", "رقم المشهد"]
    dvs = wb[px.PATTERNS_SHEET].data_validations.dataValidation
    lists = [d for d in dvs if d.type == "list"]
    assert len(lists) == sum(1 for f in pf.FIELDS if f["type"] == "select") and all(px.LISTS_SHEET in d.formula1 for d in lists)
    labels = {c.value for c in wb[px.LISTS_SHEET]["A"][1:] if c.value}
    assert {"مشكلة ← تضخيم ← حل", "نحن مقابل هم"} <= labels                    # أسماء عربية مش slugs
    readme = " ".join(str(c.value) for row in wb[px.README_SHEET].iter_rows() for c in row if c.value)
    assert "Alt+Enter" in readme and "كشف المكوّن السري" in readme and ".xlsx" in readme


def test_the_example_in_the_template_is_a_valid_pattern(lib):
    row, sc = example_pattern()
    res = px.preview(make_xlsx([row], sc), "t.xlsx")
    assert res["counts"] == {"ok": 1, "warn": 0, "error": 0} and res["rows"][0]["errors"] == []


# ═════════════════ المعاينة والتنفيذ ═════════════════
def test_full_import_converts_labels_to_values_and_records_the_source(lib):
    row, sc = example_pattern()
    prev = px.preview(make_xlsx([row], sc), "my.xlsx")
    assert prev["rows"][0]["status"] == "ok" and prev["rows"][0]["name"] == "كشف المكوّن السري"
    assert "fields" not in prev["rows"][0]                                          # الحقول ما بتنبعت للواجهة
    out = px.commit(prev["token"], by="سارة")
    assert len(out["created"]) == 1 and out["failed"] == [] and out["skipped"] == 0
    pid = out["created"][0]["id"]
    p = json.loads((lib / "library" / "patterns" / f"{pid}.pattern.json").read_text(encoding="utf-8"))
    assert pid.startswith("excel_") and p["pattern_name"] == "كشف المكوّن السري"
    assert (p["framework"], p["product_category"], p["dialect_or_language"], p["production_complexity"]) == ("pas", "fragrance", "shami", "low")
    assert p["objection_handled"] == ["هل بيثبت فعلاً؟", "هل السعر مناسب؟"] and p["hook"]["duration_sec"] == 2.5
    assert [s["function"] for s in p["scene_blueprint"]][0] == "الهوك: سؤال صادم" and len(p["scene_blueprint"]) == 3
    assert p["_source"]["type"] == "authored" and p["_source"]["via"] == "excel" and p["_source"]["file"] == "my.xlsx" and p["_source"]["row"] == 2
    assert p["_source"]["by"] == "سارة" and p.get("cluster_id")                       # العناقيد انحسبت بالآخر
    with pytest.raises(px.ExcelError, match="انتهت أو انستعملت"):
        px.commit(prev["token"])                                                    # المعاينة بتنستعمل مرة وحدة


def test_row_errors_explain_what_and_where(lib):
    good, gsc = example_pattern(pattern_name="نمط سليم")
    bad_value, bsc = example_pattern(pattern_name="قيمة غلط", framework="مش موجود بالقائمة")
    no_fw, nsc = example_pattern(pattern_name="بلا إطار")
    del no_fw["framework"]
    few, fsc = example_pattern(pattern_name="مشهدين بس")
    fsc = fsc[:2]
    prev = px.preview(make_xlsx([good, bad_value, no_fw, few], gsc + bsc + nsc + fsc), "x.xlsx")
    by = {r["name"]: r for r in prev["rows"]}
    assert prev["counts"] == {"ok": 1, "warn": 0, "error": 3}
    assert [e["field"] for e in by["قيمة غلط"]["errors"]] == ["framework"]
    assert "مش من القائمة" in by["قيمة غلط"]["errors"][0]["msg"] and by["قيمة غلط"]["errors"][0]["label"] == "الإطار"
    assert [e["label"] for e in by["بلا إطار"]["errors"]] == ["الإطار"]
    assert by["مشهدين بس"]["errors"][0]["field"] == "scene_blueprint" and "3 مشاهد" in by["مشهدين بس"]["errors"][0]["msg"]
    assert by["نمط سليم"]["row"] == 2 and by["مشهدين بس"]["row"] == 5                  # رقم السطر بالإكسل


def test_commit_saves_only_the_good_rows_and_reclusters_once(lib, monkeypatch):
    a, asc = example_pattern(pattern_name="أول")
    b, bsc = example_pattern(pattern_name="غلط", product_category="مش فئة")
    c, csc = example_pattern(pattern_name="تالت")
    calls = []
    monkeypatch.setattr(ps, "recluster", lambda: calls.append(1))
    prev = px.preview(make_xlsx([a, b, c], asc + bsc + csc))
    out = px.commit(prev["token"])
    assert [x["name"] for x in out["created"]] == ["أول", "تالت"] and out["skipped"] == 1 and calls == [1]
    assert len(ps.list_ids()) == 2


def test_commit_can_pick_rows(lib):
    a, asc = example_pattern(pattern_name="أول")
    c, csc = example_pattern(pattern_name="تالت")
    prev = px.preview(make_xlsx([a, c], asc + csc))
    out = px.commit(prev["token"], rows=[3])
    assert [x["name"] for x in out["created"]] == ["تالت"] and out["skipped"] == 1


def test_a_pattern_similar_to_the_library_is_a_warning_not_an_error(lib):
    row, sc = example_pattern()
    px.commit(px.preview(make_xlsx([row], sc))["token"])
    prev = px.preview(make_xlsx([row], sc))
    r = prev["rows"][0]
    assert r["status"] == "warn" and r["similar"][0]["same_name"] is True and "بنفس اسم نمط موجود" in r["warnings"][0]
    assert prev["counts"] == {"ok": 0, "warn": 1, "error": 0}
    assert px.commit(prev["token"])["created"] == []                                 # بدون موافقة ما بيحفظ
    prev = px.preview(make_xlsx([row], sc))
    out = px.commit(prev["token"], include_warnings=True)
    assert len(out["created"]) == 1 and len(ps.list_ids()) == 2 and out["created"][0]["id"].endswith("_2")


def test_duplicate_names_inside_one_file_are_an_error_on_the_second(lib):
    a, asc = example_pattern()
    prev = px.preview(make_xlsx([a, a], asc))
    assert [r["status"] for r in prev["rows"]] == ["ok", "error"] and "مكرّر بالملف (سطر 2)" in prev["rows"][1]["errors"][0]["msg"]


def test_notes_for_unknown_columns_orphan_scenes_and_nameless_scenes(lib):
    row, sc = example_pattern()
    orphan = {**sc[0], "اسم النمط": "نمط مش موجود"}
    nameless = {k: v for k, v in sc[0].items() if k != "اسم النمط"}

    def tweak(wb):
        wb[px.PATTERNS_SHEET].cell(row=1, column=len(pf.FIELDS) + 1, value="عمود غريب")

    prev = px.preview(make_xlsx([row], sc + [orphan, nameless], tweak=tweak))
    notes = " | ".join(prev["notes"])
    assert "عمود غريب" in notes and "نمط مش موجود" in notes and "مشهد بلا اسم نمط" in notes
    assert prev["counts"]["ok"] == 1


def test_scene_order_comes_from_the_number_column_then_the_sheet(lib):
    row, sc = example_pattern()
    prev = px.preview(make_xlsx([row], [sc[2], sc[0], sc[1]]))                       # مرتّبين غلط بالورقة، بس مرقّمين صح
    pid = px.commit(prev["token"])["created"][0]["id"]
    p = json.loads((lib / "library" / "patterns" / f"{pid}.pattern.json").read_text(encoding="utf-8"))
    assert [s["function"] for s in p["scene_blueprint"]] == [s["وظيفة المشهد"] for s in sc]


def test_lists_split_on_newlines_semicolons_and_bullets_and_digits_are_normalized(lib):
    row, sc = example_pattern(why_it_works="- سبب أول؛ سبب تاني\n• سبب تالت", **{"hook.duration_sec": "٢٫٥", "reusability_score": "٩"})
    pid = px.commit(px.preview(make_xlsx([row], sc))["token"])["created"][0]["id"]
    p = json.loads((lib / "library" / "patterns" / f"{pid}.pattern.json").read_text(encoding="utf-8"))
    assert p["why_it_works"] == ["سبب أول", "سبب تاني", "سبب تالت"] and p["hook"]["duration_sec"] == 2.5 and p["reusability_score"] == 9


def test_headers_may_drop_the_star_and_the_parenthesis_or_use_the_internal_name(lib):
    row, sc = example_pattern()

    def tweak(wb):
        ws = wb[px.PATTERNS_SHEET]
        ws.cell(row=1, column=1, value="pattern_name")
        ws.cell(row=1, column=[f["path"] for f in pf.FIELDS].index("reusability_score") + 1, value="قابلية الاستخدام")

    prev = px.preview(make_xlsx([row], sc, tweak=tweak))
    assert prev["counts"]["ok"] == 1 and prev["notes"] == []


def test_empty_sheet_and_missing_sheets_are_clear_errors(lib):
    with pytest.raises(px.ExcelError, match="ما فيها أنماط"):
        px.preview(make_xlsx([], []))

    def drop(wb):
        del wb[px.PATTERNS_SHEET]

    with pytest.raises(px.ExcelError, match="ما لقيت ورقة «الأنماط»"):
        px.preview(make_xlsx([], [], tweak=drop))


def test_too_many_patterns_in_one_file(lib):
    rows = [{"pattern_name": f"نمط {i}"} for i in range(px.MAX_PATTERNS + 1)]
    with pytest.raises(px.ExcelError, match="أكتر من 200"):
        px.preview(make_xlsx(rows, []))


# ═════════════════ حماية الملف والتوكن ═════════════════
def zip_bytes(files: dict) -> bytes:
    buf = io.BytesIO()
    with zipfile.ZipFile(buf, "w", zipfile.ZIP_DEFLATED) as z:
        for k, v in files.items():
            z.writestr(k, v)
    return buf.getvalue()


def _bad_files():
    return {
        "empty": (b"", "فاضي"),
        "old-xls": (b"\xd0\xcf\x11\xe0" + b"x" * 50, ".xls"),
        "plain-text": (b"just text, not a workbook", "مش بصيغة xlsx"),
        "broken-zip": (b"PKgarbage", "تالف"),
        "zip-without-xl": (zip_bytes({"hello.txt": "hi"}), "مش xlsx"),
        "zip-bomb": (zip_bytes({"[Content_Types].xml": "<x/>", "xl/big.bin": b"\0" * (px.MAX_UNZIPPED + 1)}), "غير طبيعي"),
        "too-big": (b"PK" + b"0" * (px.MAX_BYTES + 1), "أكبر من 5 ميغا"),
    }


@pytest.mark.parametrize("case", ["empty", "old-xls", "plain-text", "broken-zip", "zip-without-xl", "zip-bomb", "too-big"])
def test_bad_files_are_refused_with_a_reason(lib, case):
    data, msg = _bad_files()[case]
    with pytest.raises(px.ExcelError, match=msg):
        px.preview(data)


def test_a_zip_that_opens_but_is_not_a_workbook_is_refused(lib):
    with pytest.raises(px.ExcelError):
        px.preview(zip_bytes({"[Content_Types].xml": "<x/>", "xl/workbook.xml": "<not-xml"}))


def test_bad_tokens_never_touch_the_disk(lib):
    for t in ("", "../x", "zz", "0" * 15, None):
        with pytest.raises(px.ExcelError):
            px.commit(t)


def test_old_previews_are_swept(lib):
    d = px.imports_dir()
    d.mkdir(parents=True)
    old, fresh = d / ("a" * 16 + ".json"), d / ("b" * 16 + ".json")
    old.write_text("{}", encoding="utf-8")
    fresh.write_text("{}", encoding="utf-8")
    t = time.time() - px.PREVIEW_TTL - 60
    os.utime(old, (t, t))
    row, sc = example_pattern()
    px.preview(make_xlsx([row], sc))
    assert not old.exists() and fresh.exists()


def test_the_preview_lives_next_to_the_patterns_not_inside_them(lib):
    row, sc = example_pattern()
    px.preview(make_xlsx([row], sc))
    assert px.imports_dir() == lib / "library" / "pattern_imports" and ps.list_ids() == []


# ═════════════════ المسارات ═════════════════
@pytest.fixture
def client(lib):
    from fastapi.testclient import TestClient
    from web.app import app
    return TestClient(app)


def test_routes_template_preview_commit(client, lib):
    t = client.get("/api/pattern-lib/template.xlsx")
    assert t.status_code == 200 and "spreadsheetml" in t.headers["content-type"] and "attachment" in t.headers["content-disposition"]
    assert openpyxl.load_workbook(io.BytesIO(t.content)).sheetnames[0] == px.README_SHEET
    row, sc = example_pattern()
    up = client.post("/api/pattern-lib/import", files={"file": ("p.xlsx", make_xlsx([row], sc), "application/octet-stream")})
    d = up.json()
    assert up.status_code == 200 and d["counts"]["ok"] == 1 and d["filename"] == "p.xlsx"
    done = client.post("/api/pattern-lib/import/commit", json={"token": d["token"]}).json()
    assert len(done["created"]) == 1
    assert any(p["id"] == done["created"][0]["id"] for p in client.get("/api/patterns").json())
    again = client.post("/api/pattern-lib/import/commit", json={"token": d["token"]})
    assert again.status_code == 400 and "انتهت" in again.json()["detail"]


def test_routes_refuse_bad_uploads_with_400(client):
    r = client.post("/api/pattern-lib/import", files={"file": ("x.xlsx", b"not a workbook", "application/octet-stream")})
    assert r.status_code == 400 and "xlsx" in r.json()["detail"]
    assert client.post("/api/pattern-lib/import/commit", json={"token": "nope"}).status_code == 400
    assert client.post("/api/pattern-lib/import/commit", json={"token": "a" * 16, "rows": ["x"]}).status_code == 400
