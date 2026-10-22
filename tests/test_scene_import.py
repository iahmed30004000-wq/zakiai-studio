"""استيراد المشاهد من نص أو Excel (scene_import): قواعد النص، نموذج Excel، وحدود الأمان."""

import io
import os
import sys
import zipfile

import openpyxl
import pytest

sys.path.insert(0, os.path.dirname(__file__))

import scene_import as si  # noqa: E402


def test_each_paragraph_is_a_scene_and_labels_fill_fields():
    t = """ليش عيونك بتتعب آخر النهار؟
صورة: وجه امرأة بتفرك عيونها
شاشة: عيونك تعبانة؟
مدة: ٣
لقطة: قريبة
إحساس: فضول

هاي القطرة بتغيّر كل شي. جرّبها لأسبوع واحد."""
    r = si.parse_text(t)
    a, b = r["scenes"]
    assert a == {"voiceover": "ليش عيونك بتتعب آخر النهار؟", "visual": "وجه امرأة بتفرك عيونها", "on_screen_text": "عيونك تعبانة؟",
                 "duration_sec": 3.0, "shot_type": "close", "emotion": "فضول"}
    assert b["voiceover"].startswith("هاي القطرة") and b["shot_type"] is None and b["duration_sec"] == round(8 / 2.4, 1)
    assert r["warnings"] == []


def test_missing_duration_comes_from_the_words_with_a_floor_and_a_silent_default():
    one = si.parse_text("كلمة")["scenes"][0]
    assert one["duration_sec"] == si.MIN_SEC
    silent = si.parse_text("صورة: منتج على الطاولة")["scenes"][0]
    assert silent["voiceover"] == "" and silent["duration_sec"] == si.SILENT_SEC and silent["visual"] == "منتج على الطاولة"
    long = si.parse_text(" ".join(["كلمة"] * 24))["scenes"][0]
    assert long["duration_sec"] == 10.0


def test_bad_duration_and_unknown_shot_become_warnings_not_errors():
    r = si.parse_text("كلام\nمدة: abc\nلقطة: طائرة")
    assert r["scenes"][0]["shot_type"] is None
    assert len(r["warnings"]) == 2 and "المدة" in r["warnings"][0] and "نوع اللقطة" in r["warnings"][1]


def test_arabic_indic_digits_and_decimals():
    assert si.parse_text("كلام\nمدة: ٢٫٥")["scenes"][0]["duration_sec"] == 2.5
    assert si.parse_duration("4 ثواني") == 4.0 and si.parse_duration("۶") == 6.0
    assert si.parse_duration(1000) == si.MAX_DURATION_SEC
    with pytest.raises(ValueError):
        si.parse_duration("بدون رقم")
    with pytest.raises(ValueError):
        si.parse_duration(0)


def test_split_sentences_makes_one_scene_per_sentence_and_labels_go_to_the_first():
    r = si.parse_text("جملة أولى. جملة تانية؟ جملة تالتة!\nصورة: لقطة المنتج", split_sentences=True)
    assert [s["voiceover"] for s in r["scenes"]] == ["جملة أولى.", "جملة تانية؟", "جملة تالتة!"]
    assert r["scenes"][0]["visual"] == "لقطة المنتج" and r["scenes"][1]["visual"] == ""


def test_text_limits_and_empty_input():
    for bad in ("", "   \n  ", None):
        with pytest.raises(si.SceneImportError):
            si.parse_text(bad)
    with pytest.raises(si.SceneImportError) as e:
        si.parse_text("ك" * (si.MAX_TEXT + 1))
    assert "أطول" in str(e.value)
    with pytest.raises(si.SceneImportError) as e:
        si.parse_text("\n\n".join(f"مشهد {i}" for i in range(si.MAX_SCENES + 1)))
    assert str(si.MAX_SCENES) in str(e.value)


def test_a_colon_inside_normal_speech_is_not_a_label():
    r = si.parse_text("قال: لا تتأخر\nوبعدين: اطلب الآن")
    assert len(r["scenes"]) == 1 and "قال: لا تتأخر" in r["scenes"][0]["voiceover"]


def _fill(rows, header=None):
    wb = openpyxl.load_workbook(io.BytesIO(si.template()))
    ws = wb[si.SHEET_SCENES]
    if header:
        for i, h in enumerate(header, 1):
            ws.cell(row=1, column=i, value=h)
    for r in rows:
        ws.append(r)
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def test_template_has_readme_scenes_and_hidden_lists_and_roundtrips():
    wb = openpyxl.load_workbook(io.BytesIO(si.template()))
    assert wb.sheetnames == ["اقرأني", "المشاهد", "قوائم"] and wb["قوائم"].sheet_state == "hidden"
    assert wb["المشاهد"].max_row == 1                              # مثال الشرح ما بينستورد: ورقة المشاهد فاضية
    r = si.parse_excel(_fill([["الكلام الأول", "شاشة 1", "صورة 1", 3, "قريبة", "فضول"], ["مشهد صامت", None, "صورة تانية", None, "يد مع المنتج", None]]))
    assert [s["shot_type"] for s in r["scenes"]] == ["close", "hand_product"]
    assert r["scenes"][1]["duration_sec"] == si.MIN_SEC and r["scenes"][0]["on_screen_text"] == "شاشة 1"


def test_excel_unknown_columns_become_a_warning_and_empty_rows_are_skipped():
    r = si.parse_excel(_fill([["كلام", "ش", "ص", 2, "متوسطة", None, "غريب"], [None] * 7, [None, None, None, None, None, None, "فقط غريب"]],
                             header=["الكلام *", "نص الشاشة", "وصف الصورة", "المدة (ث)", "نوع اللقطة", "الإحساس", "عمود غريب"]))
    assert len(r["scenes"]) == 1 and r["warnings"] and "عمود غريب" in r["warnings"][0]


def test_excel_row_errors_are_reported_in_arabic_with_the_row_number():
    with pytest.raises(si.SceneImportError) as e:
        si.parse_excel(_fill([["كلام", None, None, "abc", None, None], ["كلام تاني", None, None, 2, "طائرة", None]]))
    msg = str(e.value)
    assert "سطر 2" in msg and "المدة" in msg and "سطر 3" in msg and "نوع اللقطة" in msg


def test_excel_limits_and_bad_files():
    rows = [[f"مشهد {i}", None, None, 2, None, None] for i in range(si.MAX_ROWS + 1)]
    with pytest.raises(si.SceneImportError) as e:
        si.parse_excel(_fill(rows))
    assert str(si.MAX_ROWS) in str(e.value)
    with pytest.raises(si.SceneImportError):
        si.parse_excel(b"")
    with pytest.raises(si.SceneImportError) as e:
        si.parse_excel(b"\xd0\xcf\x11\xe0" + b"0" * 100)
    assert "xls" in str(e.value)
    with pytest.raises(si.SceneImportError):
        si.parse_excel(b"plain text file")
    with pytest.raises(si.SceneImportError):
        si.parse_excel(b"PK" + b"0" * 100)
    zbuf = io.BytesIO()
    with zipfile.ZipFile(zbuf, "w") as z:
        z.writestr("hello.txt", "x")
    with pytest.raises(si.SceneImportError):
        si.parse_excel(zbuf.getvalue())
    with pytest.raises(si.SceneImportError) as e:
        si.parse_excel(b"PK" + b"0" * (si.MAX_BYTES + 1))
    assert "أكبر" in str(e.value)


def test_excel_without_the_scenes_sheet_uses_the_first_other_sheet_and_empty_sheet_errors():
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "ورقتي"
    ws.append(["الكلام", "المدة"])
    ws.append(["جملة", 4])
    buf = io.BytesIO()
    wb.save(buf)
    assert si.parse_excel(buf.getvalue())["scenes"][0]["duration_sec"] == 4.0
    wb2 = openpyxl.Workbook()
    wb2.active.title = "المشاهد"
    buf2 = io.BytesIO()
    wb2.save(buf2)
    with pytest.raises(si.SceneImportError):
        si.parse_excel(buf2.getvalue())
