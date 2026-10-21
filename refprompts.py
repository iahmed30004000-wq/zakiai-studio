"""
برومبتات الورقة المرجعية: للمنتج (وضعيات ثابتة) وللشخصية (وضعيات ثابتة).

الفكرة: كل وضعية (زاوية) إلها برومبت جاهز بنفس البنية بالضبط (نفس الإطار، النسب، الإضاءة، الخلفية) ولا يتغيّر فيه غير
وضعية الكاميرا — هيك الصور المولّدة (بـ ChatGPT أو Flow أو Gemini…) بتطلع ثابتة بالحجم والزوايا. الوضعية «أمام» بتتولّد أول،
وباقي الوضعيات بتنرفق معها صورة «أمام» كمرجع (قفل هوية).

البرومبتات بالإنجليزي لأن موديلات الصور أدق فيها؛ شرح كل وضعية بالعربي للواجهة. بنقل وصف المنتج/نص الملصق/اللبس كما كتبه المالك.
دوال صافية (بلا قراءة ملفات) — القراءة بـ product_ref.py وcharacters.py.
"""

from __future__ import annotations

import io
import re

# ═════════════════════════ المنتج ═════════════════════════
# ratio: أي بُعدين بيطلعوا بعرض/ارتفاع إطار المنتج المفرَّغ (للتحقق من النسبة)
PRODUCT_SLOTS = [
    {"slot": "front", "ar": "أمام", "en": "FRONT", "required": True, "ratio": ("w", "h"),
     "pose": "FRONT view: straight-on, camera at the product's mid-height and perfectly level, the product facing the camera so its main label is centered and fully readable. Long lens, no perspective distortion.",
     "pose_ar": "من قدّام بشكل مستقيم، الكاميرا بنص ارتفاع المنتج، الملصق بالنص وواضح. هي أول وضعية بتتولّد وباقي الوضعيات بتنبني عليها."},
    {"slot": "back", "ar": "خلف", "en": "BACK", "required": True, "ratio": ("w", "h"),
     "pose": "BACK view: the product rotated exactly 180° from the front so its back faces the camera; camera at mid-height, perfectly level, long lens.",
     "pose_ar": "من ورا: المنتج مدوَّر 180° بالضبط عن الأمام، بنفس الارتفاع."},
    {"slot": "left", "ar": "يسار", "en": "LEFT-SIDE", "required": True, "ratio": ("d", "h"),
     "pose": "LEFT-SIDE view: pure side profile — the product rotated exactly 90° so its LEFT side (as seen in the front view) faces the camera; camera at mid-height, perfectly level, long lens.",
     "pose_ar": "جنب يسار صافي (90°): الجنب الشمال للمنتج كما بالأمام يواجه الكاميرا."},
    {"slot": "right", "ar": "يمين", "en": "RIGHT-SIDE", "required": True, "ratio": ("d", "h"),
     "pose": "RIGHT-SIDE view: pure side profile — the product rotated exactly 90° so its RIGHT side (as seen in the front view) faces the camera; camera at mid-height, perfectly level, long lens.",
     "pose_ar": "جنب يمين صافي (90°)."},
    {"slot": "top", "ar": "فوق", "en": "TOP", "required": True, "ratio": ("w", "d"),
     "pose": "TOP view: camera directly above (90° overhead) looking straight down at the top of the product; the product's front edge points to the bottom of the frame.",
     "pose_ar": "من فوق بزاوية 90°: بنشوف راس المنتج، وواجهته باتجاه أسفل الصورة."},
    {"slot": "bottom", "ar": "تحت", "en": "BOTTOM", "required": True, "ratio": ("w", "d"),
     "pose": "BOTTOM view: camera directly below looking straight up at the base of the product (the product floats on an invisible support); the product's front edge points to the bottom of the frame.",
     "pose_ar": "من تحت: بنشوف قاعدة المنتج، وواجهته باتجاه أسفل الصورة."},
    {"slot": "in_hand", "ar": "بيد (للحجم)", "en": "SCALE-IN-HAND", "required": False, "ratio": None,
     "pose": "SCALE view: exactly one adult hand holds the product naturally at chest height, label facing the camera, fingers not covering the label, so the real size of the product relative to the hand is obvious. Camera level, medium close-up.",
     "pose_ar": "اختيارية: المنتج بيد وحدة وملصقه ظاهر، عشان حجمه الحقيقي ينثبّت بالنسبة لليد."},
    {"slot": "label", "ar": "الملصق قريب", "en": "LABEL CLOSE-UP", "required": False, "ratio": None,
     "pose": "LABEL CLOSE-UP: camera perpendicular to the label, flat and parallel, the label area filling the frame; every printed word sharp, undistorted and exactly as given.",
     "pose_ar": "اختيارية: لقطة قريبة للملصق بشكل مسطّح، عشان الكلمات تطلع صح حرفياً."},
    {"slot": "box_front", "ar": "العلبة من قدّام", "en": "BOX FRONT", "required": False, "ratio": None,
     "pose": "BOX FRONT view: the product's retail carton (outer box) alone, standing upright and photographed straight-on, camera at the box's mid-height and perfectly level, the box's front face centered and every printed word on it readable. Long lens, no perspective distortion.",
     "pose_ar": "اختيارية: كرتونة المنتج لحالها من قدّام بشكل مستقيم، ونصها المطبوع واضح."},
    {"slot": "with_box", "ar": "المنتج مع علبته", "en": "PRODUCT WITH BOX", "required": False, "ratio": None,
     "pose": "PRODUCT WITH ITS BOX: the product standing beside its retail carton, both fully visible and not overlapping, each at its real relative scale, the box front and the product's label both facing the camera; camera at mid-height and perfectly level.",
     "pose_ar": "اختيارية: المنتج واقف جنب كرتونته، الاثنين ظاهرين وبحجمهم النسبي الحقيقي."},
]
PRODUCT_SLOT_KEYS = [s["slot"] for s in PRODUCT_SLOTS]
BOX_SLOTS = ("box_front", "with_box")                 # وضعيات الكرتونة (اختيارية، بلا فحص نسبة)
_PSPEC = {s["slot"]: s for s in PRODUCT_SLOTS}

_FRAMING = ("Framing (identical for every view of this product): square 1:1 canvas, the product centered and completely inside the frame, "
            "occupying 80% of the image height with equal margins, no cropping.")
_LIGHT = ("Lighting & background (identical for every view): soft even studio light from the front-left, only a very soft contact shadow, "
          "plain pure-white seamless background.")


def _num(v) -> str:
    return ("%g" % round(float(v), 2))


def product_prompt(slot: str, *, name: str = "", category: str = "", visual: str = "", dims: dict | None = None,
                   label: str = "", colors: list | None = None, box_text: str = "") -> str:
    spec = _PSPEC[slot]
    dims = dims or {}
    box_only = slot == "box_front"                        # الكرتونة لحالها: أبعاد المنتج ونص ملصقه ما بيخصّوها
    lines = [f"PRODUCT REFERENCE SHEET — {spec['en']} VIEW. Photorealistic studio product photo (not an illustration, not a 3D render)."]
    who = (name or "the product").strip()
    cat = re.sub(r"[_-]+", " ", category or "").strip()
    lines.append(f"Product: {who}{f' ({cat})' if cat else ''}." + (f' Look (as defined by the owner): "{visual.strip()}".' if visual and visual.strip() else ""))
    if box_only:
        lines.append("Keep the box's proportions realistic for a retail carton that holds this product. Never stretch, squash or redesign the box.")
    elif all(dims.get(k) for k in ("h", "w", "d")):
        h, w, d = (_num(dims[k]) for k in ("h", "w", "d"))
        lines.append(f"True proportions: height {h} cm, width {w} cm, depth {d} cm (H:W:D = {h}:{w}:{d}). Never stretch, squash or redesign the product.")
    else:
        lines.append("Keep the product's proportions identical in every view. Never stretch, squash or redesign the product.")
    cols = [str(c).strip() for c in (colors or []) if str(c).strip()]
    if cols:
        lines.append("Colors: " + ", ".join(cols) + ".")
    if not box_only:
        if label and label.strip():
            lines.append(f'Label text — reproduce it letter for letter, with no extra or invented words: "{label.strip()}". Arabic text must stay right-to-left, sharp and undistorted.')
        else:
            lines.append("Do not add any text to the product.")
    if slot in BOX_SLOTS:
        if box_text and box_text.strip():
            lines.append(f'Box text — reproduce it letter for letter, with no extra or invented words: "{box_text.strip()}". Arabic text must stay right-to-left, sharp and undistorted.')
        else:
            lines.append("Do not add any text to the box.")
    if box_only:
        lines.append("IDENTITY LOCK: use the attached FRONT-view image of the product to match its brand look, colors and graphics on the box. "
                     "Only the object shown changes: here it is the retail box.")
    elif slot != "front":
        lines.append("IDENTITY LOCK: use the attached FRONT-view image as the exact same product — identical shape, materials, colors, graphics, "
                     "label layout and proportions. Only the viewing angle changes.")
    lines.append("Camera/pose: " + spec["pose"])
    lines.append("Framing: the label area fills the frame, no cropping of the label." if slot == "label" else _FRAMING)
    lines.append(_LIGHT)
    if slot == "in_hand":
        lines.append("Rules: one product and exactly one hand only; no other props, people, extra text, logos, watermarks or reflections; sharp focus, true-to-life colors.")
    elif box_only:
        lines.append("Rules: one retail box only; no product bottle, hands, props, people, extra text, logos, watermarks or reflections of other objects; sharp focus, true-to-life colors.")
    elif slot == "with_box":
        lines.append("Rules: one product and its box only; no hands, props, people, extra text, logos, watermarks or reflections of other objects; sharp focus, true-to-life colors.")
    else:
        lines.append("Rules: one product only; no hands, props, people, extra text, logos, watermarks or reflections of other objects; sharp focus, true-to-life colors.")
    return "\n".join(lines)


def product_gaps(*, dims: dict | None, label: str, colors: list | None, visual: str) -> list[str]:
    """شو ناقص ويضعّف ثبات البرومبتات (بالعربي، للواجهة)."""
    out = []
    if not all((dims or {}).get(k) for k in ("h", "w", "d")):
        out.append("الأبعاد بالسنتيمتر (بتثبّت النِّسب والحجم)")
    if not (visual or "").strip():
        out.append("وصف شكل المنتج")
    if not (label or "").strip():
        out.append("نص الملصق الحرفي")
    if not [c for c in (colors or []) if str(c).strip()]:
        out.append("الألوان")
    return out


# ═════════════════════════ الشخصية ═════════════════════════
CHARACTER_SLOTS = [
    {"slot": "front", "ar": "أمام", "en": "FRONT", "required": True,
     "pose": "Head-and-shoulders portrait, facing the camera straight-on with eyes to the lens, relaxed friendly neutral expression, camera at eye level (85mm look).",
     "pose_ar": "وجه وكتفين من قدّام والنظر للكاميرا. أول وضعية بتتولّد وباقي الوضعيات بتنبني عليها."},
    {"slot": "side", "ar": "جانب", "en": "SIDE PROFILE", "required": True,
     "pose": "Head-and-shoulders portrait in pure side profile (90°), looking straight ahead, same expression, camera at eye level (85mm look).",
     "pose_ar": "وجه وكتفين من الجنب (90°)."},
    {"slot": "three_quarter", "ar": "ثلاثة أرباع", "en": "THREE-QUARTER", "required": True,
     "pose": "Waist-up portrait, the body turned 45° from the camera and the face turned toward the lens (three-quarter view), camera at eye level.",
     "pose_ar": "من الوسط لفوق، الجسم مايل 45° والوجه نحو الكاميرا."},
    {"slot": "full_body", "ar": "جسم كامل", "en": "FULL BODY", "required": True,
     "pose": "Full-body shot, standing relaxed and facing the camera, the entire body from head to feet inside the frame with the feet visible, camera at chest height.",
     "pose_ar": "الجسم كامل واقف، من الراس للقدمين."},
    {"slot": "hands", "ar": "يدين", "en": "HANDS", "required": True,
     "pose": "Close-up of both hands, palms and backs shown naturally, relaxed and slightly open as if about to hold a small product; forearms in the outfit's sleeves; the same skin tone as the person's face.",
     "pose_ar": "لقطة قريبة لليدين بنفس لون البشرة، عشان لقطات حمل المنتج."},
]
CHARACTER_SLOT_KEYS = [s["slot"] for s in CHARACTER_SLOTS]
_CSPEC = {s["slot"]: s for s in CHARACTER_SLOTS}


def character_prompt(view: str, ch: dict, country: str = "") -> str:
    spec = _CSPEC[view]
    male = ch.get("gender") == "m"
    who = "man" if male else "woman"
    age = ch.get("age") or ""
    lines = [f"CHARACTER REFERENCE SHEET — {spec['en']}. Photorealistic studio photo of the SAME person in every view of this sheet."]
    persona = str(ch.get("persona") or "").strip()
    lines.append(f"Person: {age}-year-old {who}" + (f", {persona}" if persona else "") + ".")
    attrs = []
    if str(ch.get("skin") or "").strip():
        attrs.append(f"Skin: {ch['skin'].strip()}")
    if str(ch.get("hair") or "").strip():
        attrs.append(f"Hair: {ch['hair'].strip()}")
    looks = ch.get("market_look") or {}
    if looks:
        pick = (country or "").upper()
        if pick in looks:
            text = looks[pick]
        elif len(looks) == 1:
            text = next(iter(looks.values()))
        else:
            text = "; ".join(f"{c}: {t}" for c, t in list(looks.items())[:2])
        if str(text or "").strip():
            attrs.append(f"Regional look: {str(text).strip()}")
    if attrs:
        lines.append(". ".join(attrs) + ".")
    ward = ch.get("wardrobe") or {}
    if str(ward.get("description") or "").strip():
        pal = [str(c).strip() for c in (ward.get("palette") or []) if str(c).strip()]
        lines.append(f"Outfit (identical in every view): {ward['description'].strip()}." + (f" Palette: {', '.join(pal)}." if pal else ""))
    if ch.get("source") == "real_person":
        lines.append("IDENTITY SOURCE: the attached photos of the real, consenting person are the ONLY source of identity — do not invent or change any facial feature.")
    if view != "front":
        lines.append("IDENTITY LOCK: use the attached FRONT image as the exact same person — identical face and facial structure, skin tone, hairstyle, age and outfit. "
                     "Do not redesign anything; only the pose and camera angle change.")
    lines.append("Pose: " + spec["pose"])
    lines.append("Background & light (identical for every view): plain light-gray seamless background, soft even studio lighting, no shadow on the background.")
    lines.append("Rules: one person only; no text, logos, watermarks, props or jewelry changes; natural skin texture, no beauty filters or retouching that alters identity.")
    return "\n".join(lines)


def character_gaps(ch: dict) -> list[str]:
    out = []
    for k, label in (("skin", "البشرة"), ("hair", "الشعر"), ("persona", "الطابع")):
        if not str(ch.get(k) or "").strip():
            out.append(label)
    if not str((ch.get("wardrobe") or {}).get("description") or "").strip():
        out.append("وصف اللبس")
    if not ch.get("age"):
        out.append("العمر")
    return out


# ═════════════════════════ شو بيوصل للتوليد فعلياً ═════════════════════════
# image_studio.refs_with_roles بيقرا من هون (مصدر واحد) والواجهة بتعرض الشارة منهم، فما بيتفرّقوا.
CHARACTER_FACE_VIEWS = ("front", "three_quarter")       # بتنبعت مع لقطات الوجه
CHARACTER_HAND_VIEWS = ("hands",)                       # بتنبعت مع لقطات اليد بالمنتج
PRODUCT_SENT_SLOT = "front"                             # بتنبعت مع كل لقطة وبتنركّب بالصور القريبة


def product_use_note(slot: str) -> str:
    return "بتنبعت مع كل لقطة وبتنركّب بالصور القريبة" if slot == PRODUCT_SENT_SLOT else ""


def character_use_note(view: str) -> str:
    if view in CHARACTER_FACE_VIEWS:
        return "بتنبعت مع لقطات الوجه"
    if view in CHARACTER_HAND_VIEWS:
        return "بتنبعت مع لقطات اليد بالمنتج"
    return ""


def image_flag(has_image: bool, locked: bool) -> str:
    """حالة صورة الوضعية بالـ Word: مثبّتة / موجودة / ناقصة."""
    return "الصورة: مثبّتة" if locked else ("الصورة: موجودة" if has_image else "الصورة: ناقصة")


def slot_docx_items(slots: list[dict]) -> list[dict]:
    """بنود Word من وضعيات (نفس شكل product_ref.sheet / characters.slots): العنوان وشرح الوضعية وشارات
    (حالة الصورة، «ارفق صورة أمام»، معدَّل، بتنبعت للتوليد) والبرومبت."""
    out = []
    for s in slots:
        flags = [image_flag(s["has_image"], s["locked"])]
        if s["needs_front"]:
            flags.append("ارفق صورة «أمام» معه")
        if s["prompt"]["edited"]:
            flags.append("برومبت معدَّل")
        if s.get("use_note"):
            flags.append(s["use_note"])
        out.append({"title": f"{s['label']} — {s['label_en']}" + ("" if s["required"] else " (اختيارية)"), "note": s["pose_ar"],
                    "flags": flags, "prompt": s["prompt"]["text"]})
    return out


# ═════════════════════════ Word ═════════════════════════
DOCX_MIME = "application/vnd.openxmlformats-officedocument.wordprocessingml.document"


def build_docx(title: str, intro: list[str], items: list[dict]) -> bytes:
    """items: [{title, note, prompt, flags:[str]}] — عنوان عربي، شرح عربي، برومبت إنجليزي (LTR). بيرجّع bytes لملف docx."""
    from docx import Document
    from docx.enum.text import WD_BREAK
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    from docx.shared import Pt

    def font(run, name="Arial", size=None, bold=None):
        run.font.name = name
        rpr = run._r.get_or_add_rPr()
        rfonts = rpr.find(qn("w:rFonts"))
        if rfonts is None:
            rfonts = OxmlElement("w:rFonts")
            rpr.append(rfonts)
        rfonts.set(qn("w:cs"), name)
        if size:
            run.font.size = Pt(size)
        if bold is not None:
            run.font.bold = bold

    def rtl(par):
        ppr = par._p.get_or_add_pPr()
        ppr.append(OxmlElement("w:bidi"))
        for r in par.runs:
            r._r.get_or_add_rPr().append(OxmlElement("w:rtl"))

    doc = Document()
    p = doc.add_paragraph()
    r = p.add_run(title)
    font(r, size=18, bold=True)
    rtl(p)
    for line in intro:
        p = doc.add_paragraph()
        font(p.add_run(line), size=11)
        rtl(p)
    n = 0                                                       # رقم الوضعية داخل القسم الحالي (بيبدأ من 1 مع كل عنوان فاصل)
    sections = 0
    for i, it in enumerate(items):
        if it.get("heading"):                                   # عنوان فاصل (مثلاً اسم منتج) بلا برومبت
            if sections:
                doc.add_page_break()
            sections += 1
            n = 0
            p = doc.add_paragraph()
            font(p.add_run(it["title"]), size=16, bold=True)
            rtl(p)
            for line in it.get("lines", []):
                p = doc.add_paragraph()
                font(p.add_run(line), size=11)
                rtl(p)
            continue
        n += 1
        p = doc.add_paragraph()
        font(p.add_run(f"{n}. {it['title']}"), size=14, bold=True)
        rtl(p)
        if it.get("flags"):
            p = doc.add_paragraph()
            font(p.add_run(" · ".join(it["flags"])), size=10)
            rtl(p)
        if it.get("note"):
            p = doc.add_paragraph()
            font(p.add_run(it["note"]), size=11)
            rtl(p)
        for ln in str(it["prompt"]).split("\n"):
            p = doc.add_paragraph()
            font(p.add_run(ln), name="Consolas", size=10)
            p.paragraph_format.space_after = Pt(2)
        if i + 1 < len(items) and not items[i + 1].get("heading"):
            doc.add_paragraph().add_run().add_break(WD_BREAK.LINE)
    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()
