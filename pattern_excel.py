"""
استيراد أنماط من Excel (xlsx): نموذج بتنزّله، معاينة لكل سطر قبل ما ينحفظ شي، وبعدين «استورد الصح».

  template()                        → bytes (3 أوراق + ورقة مخفية للقوائم المنسدلة)
  preview(data, filename)           → {token, filename, rows[…], counts{…}, notes[…]}  (بدون أي كتابة بالمكتبة)
  commit(token, …)                  → {created[…], failed[…], skipped}  (create لكل سطر صح، وعناقيد مرة وحدة بالآخر)

  · ورقة «الأنماط»: سطر لكل نمط. ورقة «المشاهد»: مشاهد كل نمط تحت اسمه (3 على الأقل).
  · القيم المغلقة (الإطار، الفئة، اللهجة…) بتنكتب بالاسم العربي من قائمة منسدلة، وبتتحوّل للقيمة الداخلية.
  · الملف: xlsx بس، ≤ 5 ميغا. المعاينة بتنحفظ بمجلد pattern_imports/ جنب مجلد الأنماط وبتنحذف بعد 24 ساعة.
"""

from __future__ import annotations

import io
import json
import re
import secrets
import time
import zipfile
from pathlib import Path

import config
import pattern_fields as pf
import pattern_store
import taxonomy

README_SHEET, PATTERNS_SHEET, SCENES_SHEET, LISTS_SHEET = "اقرأني", "الأنماط", "المشاهد", "قوائم"
MAX_BYTES = 5 * 1024 * 1024
MAX_UNZIPPED = 50 * 1024 * 1024
MAX_PATTERNS = 200
VALIDATION_ROWS = 500
PREVIEW_TTL = 24 * 3600
TOKEN_RX = re.compile(r"^[0-9a-f]{16}$")

SCENE_NAME_COL, SCENE_NO_COL = "اسم النمط", "رقم المشهد"
_LIST_SPLIT = re.compile(r"[\n\r؛;|]+")
_BULLET = re.compile(r"^\s*[-•*·]+\s*")


class ExcelError(Exception):
    """رسالة عربية جاهزة للعرض (400)."""


# ═════════════════ أدوات ═════════════════
def imports_dir() -> Path:
    return Path(config.PATTERNS_DIR).parent / "pattern_imports"


def _norm(s) -> str:
    """مفتاح مقارنة للعناوين والقيم: بدون نجمة المطلوب ومسافات زيادة."""
    s = re.sub(r"[\s*:：]+", " ", str(s or "")).strip().casefold()
    return s


def _plain(label: str) -> str:
    """العنوان بدون التوضيح اللي بين قوسين: «قابلية الاستخدام (0–10)» ← «قابلية الاستخدام»."""
    return _norm(re.sub(r"\s*\([^)]*\)\s*$", "", str(label or "")))


def _options() -> dict:
    """{اسم القائمة: [(slug, تسمية عربية)]} بنفس ترتيب الواجهة."""
    o = pf.options()
    return {k: list(v.items()) for k, v in o.items()}


def _select_maps() -> dict:
    """{اسم القائمة: {تسمية/slug مطبّع: slug}}."""
    out = {}
    for name, items in _options().items():
        m = {}
        for slug, label in items:
            m[_norm(label)] = slug
            m[_norm(slug)] = slug
        out[name] = m
    return out


def _header(spec: dict) -> str:
    return spec["label"] + (" *" if spec.get("required") else "")


def _scene_cols() -> list[dict]:
    return [{"key": SCENE_NAME_COL, "label": SCENE_NAME_COL + " *", "width": 28},
            {"key": SCENE_NO_COL, "label": SCENE_NO_COL, "width": 12}] + [
        {"key": f["key"], "label": f["label"] + (" *" if f.get("required") else ""), "width": 34 if f["type"] == "text" else 14}
        for f in pf.SCENE_FIELDS]


# ═════════════════ النموذج ═════════════════
def _example() -> tuple[dict, list[dict]]:
    """نمط مثال كامل (تسميات عربية متل ما بتنكتب بالخلايا)، بيطلع بورقة «اقرأني»، وبيستعمله اختبار الاستيراد."""
    fw, cat, dia = taxonomy.FRAMEWORKS["pas"], taxonomy.CATEGORY_LABELS["fragrance"], taxonomy.DIALECTS["shami"]
    fields = {
        "pattern_name": "كشف المكوّن السري",
        "one_line_summary": "افتتاحية بسؤال بيخلّي المشاهد يكمّل، ثم كشف بطيء للمكوّن اللي بيميّز المنتج.",
        "framework": fw, "product_category": cat, "dialect_or_language": dia,
        "target_audience": "نساء 25–40 بيدوّروا على عطر بيثبت طول اليوم",
        "reusability_score": 8, "production_complexity": pf.COMPLEXITY["low"],
        "hook.type": "سؤال صادم", "hook.duration_sec": 2.5,
        "hook.mechanism": "السؤال بيوقف التمرير لأنو بيلمس مشكلة المشاهد مباشرة.",
        "hook.transferable_template": "ليش [المشكلة] وما حدا بيحكيلك السبب؟",
        "product_entry_sec": 6, "product_entry_style": "المنتج بيدخل بإيد الشخصية وهي بتحل المشكلة.",
        "proof_mechanism": "كشف المكوّنات",
        "objection_handled": ["هل بيثبت فعلاً؟", "هل السعر مناسب؟"],
        "cta.type": "طلب مباشر", "cta.placement_sec": 28, "cta.wording_style": "جملة قصيرة + رقم الطلب عالشاشة",
        "pacing.total_duration_sec": 30, "pacing.cut_rhythm": "قطع سريع بالبداية وبطيء عند الكشف", "pacing.words_per_second": 2.4,
        "visual_language.shot_types": ["لقطة قريبة", "لقطة ماكرو"], "visual_language.setting": "طاولة داكنة بإضاءة دافئة",
        "visual_language.production_style": "استوديو بسيط", "visual_language.text_overlay_style": "عناوين عريضة",
        "visual_language.talent": "بنت بتحكي للكاميرا",
        "emotional_arc": ["فضول", "ترقّب", "ثقة"],
        "why_it_works": ["السؤال بيخلق فجوة معرفة", "الكشف البطيء بيبني الثقة"],
        "weak_points": ["بيحتاج مكوّن مميز يستاهل الكشف"], "best_for": ["العطور", "العناية بالبشرة"], "avoid_for": ["الأدوات المنزلية"],
    }
    scenes = [
        {"function": "الهوك: سؤال صادم", "duration_sec": 3, "shot_type": "لقطة قريبة", "vo_intent": "سؤال", "on_screen_text_role": "تكرار السؤال", "emotion": "فضول"},
        {"function": "كشف المكوّن ببطء", "duration_sec": 20, "shot_type": "لقطة ماكرو", "vo_intent": "شرح", "on_screen_text_role": "اسم المكوّن", "emotion": "ترقّب"},
        {"function": "الدعوة للشراء", "duration_sec": 7, "shot_type": "لقطة ثابتة", "vo_intent": "طلب", "on_screen_text_role": "رقم الطلب", "emotion": "ثقة"},
    ]
    return fields, scenes


def template() -> bytes:
    """ملف النموذج (xlsx)."""
    from openpyxl import Workbook
    from openpyxl.comments import Comment
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter
    from openpyxl.worksheet.datavalidation import DataValidation

    wb = Workbook()
    head_fill = PatternFill("solid", fgColor="1F4E4A")
    head_font = Font(bold=True, color="FFFFFF", size=11)
    wrap = Alignment(wrap_text=True, vertical="top", horizontal="right")
    opts = _options()

    # ── ورقة «اقرأني» ──
    ws = wb.active
    ws.title = README_SHEET
    ws.sheet_view.rightToLeft = True
    ws.column_dimensions["A"].width = 34
    ws.column_dimensions["B"].width = 90
    lines = [
        ("نموذج استيراد الأنماط — Zakiai Systems", True),
        ("", False),
        ("1) ورقة «الأنماط»: سطر لكل نمط. الأعمدة اللي عليها * مطلوبة.", False),
        ("2) ورقة «المشاهد»: كل مشهد بسطر. اكتب اسم النمط بنفس الشكل بالظبط، ورقّم المشاهد (3 مشاهد على الأقل لكل نمط).", False),
        ("3) الأعمدة اللي فيها قائمة منسدلة (الإطار، الفئة، اللهجة…): اختر منها، ما تكتب من راسك.", False),
        ("4) الأعمدة اللي فيها قائمة (مسار المشاعر، ليش بيشتغل…): كل بند بسطر لحاله داخل الخلية (Alt+Enter) أو افصل بينهم بـ ؛", False),
        ("5) احفظ الملف بصيغة .xlsx وارفعه من صفحة الأنماط ← «استورد من Excel».", False),
        ("6) قبل ما ينحفظ أي شي بتشوف معاينة لكل سطر (صح أو غلط وليش). السطور الغلط ما بتنحفظ، وبتصلّحها وبترفع الملف من جديد.", False),
        ("", False),
        ("مثال معبّى (نفس هالقيم بتنكتب بالورقتين):", True),
    ]
    for i, (text, bold) in enumerate(lines, 1):
        c = ws.cell(row=i, column=1, value=text)
        c.font = Font(bold=bold, size=13 if i == 1 else 11)
        c.alignment = Alignment(horizontal="right")
        ws.merge_cells(start_row=i, start_column=1, end_row=i, end_column=2)
    ex_fields, ex_scenes = _example()
    r = len(lines) + 1
    for spec in pf.FIELDS:
        v = ex_fields.get(spec["path"])
        if v is None:
            continue
        ws.cell(row=r, column=1, value=_header(spec)).font = Font(bold=True)
        cell = ws.cell(row=r, column=2, value="\n".join(v) if isinstance(v, list) else v)
        cell.alignment = wrap
        r += 1
    r += 1
    ws.cell(row=r, column=1, value="مشاهد المثال").font = Font(bold=True)
    for i, sc in enumerate(ex_scenes, 1):
        r += 1
        ws.cell(row=r, column=1, value=f"مشهد {i}")
        ws.cell(row=r, column=2, value=" — ".join(str(sc[k]) for k in ("function", "duration_sec", "shot_type") if sc.get(k) is not None)).alignment = wrap

    # ── ورقة «قوائم» (مخفية): مصدر القوائم المنسدلة ──
    lists = wb.create_sheet(LISTS_SHEET)
    names = ["frameworks", "categories", "dialects", "dialect_subs", "complexity"]
    for col, name in enumerate(names, 1):
        lists.cell(row=1, column=col, value=name)
        for row, (_slug, label) in enumerate(opts[name], 2):
            lists.cell(row=row, column=col, value=label)
    lists.sheet_state = "hidden"

    def list_formula(name: str) -> str:
        col = get_column_letter(names.index(name) + 1)
        return f"='{LISTS_SHEET}'!${col}$2:${col}${len(opts[name]) + 1}"

    # ── ورقة «الأنماط» ──
    wp = wb.create_sheet(PATTERNS_SHEET, 1)
    wp.sheet_view.rightToLeft = True
    for col, spec in enumerate(pf.FIELDS, 1):
        c = wp.cell(row=1, column=col, value=_header(spec))
        c.fill, c.font, c.alignment = head_fill, head_font, Alignment(wrap_text=True, vertical="center", horizontal="center")
        if spec.get("hint") or spec["type"] == "list":
            hint = (spec.get("hint") or "") + ("\nكل بند بسطر لحاله (Alt+Enter) أو افصل بـ ؛" if spec["type"] == "list" else "")
            c.comment = Comment(hint.strip(), "Zakiai")
        letter = get_column_letter(col)
        wp.column_dimensions[letter].width = {"textarea": 40, "list": 34, "text": 26, "select": 22}.get(spec["type"], 14)
        rng = f"{letter}2:{letter}{VALIDATION_ROWS + 1}"
        if spec["type"] == "select":
            dv = DataValidation(type="list", formula1=list_formula(spec["options"]), allow_blank=True, showErrorMessage=True,
                                errorTitle="قيمة غير صحيحة", error="اختر من القائمة المنسدلة")
            wp.add_data_validation(dv)
            dv.add(rng)
        elif spec["type"] in ("number", "int"):
            dv = DataValidation(type="whole" if spec["type"] == "int" else "decimal", operator="between",
                                formula1=str(spec.get("min", 0)), formula2=str(spec.get("max", 600)), allow_blank=True,
                                showErrorMessage=True, errorTitle="رقم غير صحيح", error=f"رقم بين {spec.get('min', 0)} و{spec.get('max', 600)}")
            wp.add_data_validation(dv)
            dv.add(rng)
        if spec["type"] in ("textarea", "list"):
            for row in range(2, 12):
                wp.cell(row=row, column=col).alignment = wrap
    wp.row_dimensions[1].height = 36
    wp.freeze_panes = "B2"

    # ── ورقة «المشاهد» ──
    wsn = wb.create_sheet(SCENES_SHEET, 2)
    wsn.sheet_view.rightToLeft = True
    for col, spec in enumerate(_scene_cols(), 1):
        c = wsn.cell(row=1, column=col, value=spec["label"])
        c.fill, c.font, c.alignment = head_fill, head_font, Alignment(wrap_text=True, vertical="center", horizontal="center")
        wsn.column_dimensions[get_column_letter(col)].width = spec["width"]
    wsn.row_dimensions[1].height = 30
    wsn.freeze_panes = "A2"

    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


# ═════════════════ القراءة ═════════════════
def _check_file(data: bytes) -> None:
    if not data:
        raise ExcelError("الملف فاضي")
    if len(data) > MAX_BYTES:
        raise ExcelError(f"الملف أكبر من {MAX_BYTES // (1024 * 1024)} ميغا")
    if data[:4] == b"\xd0\xcf\x11\xe0":
        raise ExcelError("هذا ملف Excel قديم (.xls) — افتحه واحفظه كـ .xlsx وارفعه من جديد")
    if data[:2] != b"PK":
        raise ExcelError("الملف مش بصيغة xlsx — نزّل النموذج وعبّيه")
    try:
        zf = zipfile.ZipFile(io.BytesIO(data))
        infos = zf.infolist()
        names = set(zf.namelist())
    except zipfile.BadZipFile:
        raise ExcelError("الملف تالف أو مش xlsx")
    if len(infos) > 300 or sum(i.file_size for i in infos) > MAX_UNZIPPED:
        raise ExcelError("الملف كبير بشكل غير طبيعي")
    if "[Content_Types].xml" not in names or not any(n.startswith("xl/") for n in names):
        raise ExcelError("الملف مش xlsx — نزّل النموذج وعبّيه")


def _cell(v):
    """قيمة الخلية: نص مقصوص، أو رقم (3.0 ← 3)، أو None للفاضي."""
    if v is None:
        return None
    if isinstance(v, bool):
        return str(v)
    if isinstance(v, float) and v == int(v):
        return int(v)
    if isinstance(v, (int, float)):
        return v
    s = str(v).strip()
    return s or None


def _split_list(v) -> list[str]:
    if v is None:
        return []
    parts = _LIST_SPLIT.split(str(v))
    return [p for p in (_BULLET.sub("", x).strip() for x in parts) if p]


def _read_sheet(wb, names: tuple[str, ...]):
    for n in wb.sheetnames:
        if _norm(n) in {_norm(x) for x in names}:
            return wb[n]
    return None


def _header_map(header_row, specs: list[dict], label_of) -> tuple[dict, list[str]]:
    """{فهرس العمود: spec} + أسماء أعمدة غير معروفة (بتنتجاهل بتحذير)."""
    lookup = {}
    for s in specs:
        lookup[_norm(label_of(s))] = s
        lookup[_plain(label_of(s))] = s
        lookup[_norm(s.get("path") or s.get("key"))] = s
    mapped, unknown = {}, []
    for i, h in enumerate(header_row):
        if h is None or not str(h).strip():
            continue
        s = lookup.get(_norm(h)) or lookup.get(_plain(h))
        if s is None:
            unknown.append(str(h).strip())
        elif all(s is not m for m in mapped.values()):
            mapped[i] = s
    return mapped, unknown


def _parse_workbook(data: bytes):
    import openpyxl
    _check_file(data)
    try:
        wb = openpyxl.load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    except Exception as e:                                          # noqa: BLE001 — أي خلل بالقراءة = ملف مش صالح
        raise ExcelError("ما قدرت أقرأ الملف — تأكد إنو xlsx سليم") from e
    wsp = _read_sheet(wb, (PATTERNS_SHEET, "patterns"))
    if wsp is None:
        raise ExcelError(f"ما لقيت ورقة «{PATTERNS_SHEET}» — استعمل النموذج (زر «نزّل النموذج»)")
    wss = _read_sheet(wb, (SCENES_SHEET, "scenes"))

    notes: list[str] = []
    maps = _select_maps()
    rows = list(wsp.iter_rows(values_only=True))
    if not rows:
        raise ExcelError(f"ورقة «{PATTERNS_SHEET}» فاضية")
    colmap, unknown = _header_map(rows[0], pf.FIELDS, _header)
    if unknown:
        notes.append("أعمدة مش معروفة وانتجاهلت: " + "، ".join(unknown))
    if not any(s["path"] == "pattern_name" for s in colmap.values()):
        raise ExcelError("ما لقيت عمود «اسم النمط» — استعمل النموذج")

    patterns = []
    for rn, row in enumerate(rows[1:], 2):
        vals = {i: _cell(row[i]) if i < len(row) else None for i in colmap}
        if not any(v is not None for v in vals.values()):
            continue
        fields, errs = {}, []
        for i, spec in colmap.items():
            v = vals[i]
            kind, path = spec["type"], spec["path"]
            if v is None:
                continue
            if kind == "list":
                pf.set_path(fields, path, _split_list(v))
            elif kind == "select":
                slug = maps[spec["options"]].get(_norm(v))
                if slug is None and path == "production_complexity":
                    slug = pf.normalize_complexity(v)
                if slug is None:
                    errs.append({"field": path, "label": spec["label"], "msg": f"«{v}» مش من القائمة — اختر من القائمة المنسدلة"})
                else:
                    pf.set_path(fields, path, slug)
            else:
                pf.set_path(fields, path, v)
        patterns.append({"row": rn, "fields": fields, "errors": errs})
        if len(patterns) > MAX_PATTERNS:
            raise ExcelError(f"أكتر من {MAX_PATTERNS} نمط بالملف — قسّمه لملفين")

    scenes_by_name: dict[str, list[dict]] = {}
    scene_rows = list(wss.iter_rows(values_only=True)) if wss is not None else []
    if scene_rows:
        scols = _scene_cols()
        sspecs = [{"key": c["key"], "label": c["label"]} for c in scols]
        smap, sunknown = _header_map(scene_rows[0], sspecs, lambda s: s["label"])
        if sunknown:
            notes.append("أعمدة مشاهد مش معروفة وانتجاهلت: " + "، ".join(sunknown))
        for rn, row in enumerate(scene_rows[1:], 2):
            vals = {smap[i]["key"]: (_cell(row[i]) if i < len(row) else None) for i in smap}
            name = vals.pop(SCENE_NAME_COL, None)
            number = vals.pop(SCENE_NO_COL, None)
            if not any(v is not None for v in vals.values()) and name is None:
                continue
            if name is None:
                notes.append(f"مشهد بلا اسم نمط (سطر {rn} بورقة المشاهد) — انتجاهل")
                continue
            sc = {k: v for k, v in vals.items() if v is not None}
            sc["_order"] = number if isinstance(number, (int, float)) else None
            sc["_row"] = rn
            scenes_by_name.setdefault(_norm(name), []).append(sc)
    wb.close()
    return patterns, scenes_by_name, notes


def _field_label(path: str) -> str:
    spec = pf.FIELD_BY_PATH.get(path)
    if spec:
        return spec["label"]
    m = re.match(r"scene_blueprint\.(\d+)(?:\.(\w+))?$", path)
    if m:
        sl = {f["key"]: f["label"] for f in pf.SCENE_FIELDS}.get(m.group(2) or "", "")
        return f"مشهد {m.group(1)}" + (f" — {sl}" if sl else "")
    return "المشاهد" if path == "scene_blueprint" else path


def preview(data: bytes, filename: str = "") -> dict:
    """يقرأ ويتحقق ويقارن مع المكتبة، وبيحفظ المعاينة بتوكن. ما بيكتب شي بالمكتبة."""
    patterns, scenes_by_name, notes = _parse_workbook(data)
    seen: dict[str, int] = {}
    used_scene_keys = set()
    rows = []
    for p in patterns:
        fields, errs = p["fields"], list(p["errors"])
        name = str(fields.get("pattern_name") or "").strip()
        key = _norm(name)
        sc = scenes_by_name.get(key, []) if key else []
        used_scene_keys.add(key)
        sc = sorted(sc, key=lambda s: (s["_order"] is None, s["_order"] if s["_order"] is not None else 0, s["_row"]))
        fields = {**fields, "scene_blueprint": [{k: v for k, v in s.items() if not k.startswith("_")} for s in sc]}
        _, verrs = pf.validate(fields, "create")
        have = {e["field"] for e in errs}                       # القيمة الغلط بالقائمة ما بتنعدّ مرتين («مش من القائمة» + «مطلوب»)
        for e in verrs:
            if e["field"] not in have:
                errs.append({"field": e["field"], "label": _field_label(e["field"]), "msg": e["msg"]})
        if key and key in seen:
            errs.append({"field": "pattern_name", "label": "اسم النمط", "msg": f"الاسم مكرّر بالملف (سطر {seen[key]})"})
        elif key:
            seen[key] = p["row"]
        warnings, similar = [], []
        if not errs:
            for m in pattern_store.similar(fields):
                if m.get("is_duplicate") or m.get("same_name"):
                    similar.append({"pattern_id": m["pattern_id"], "pattern_name": m["pattern_name"], "score": m["score"], "same_name": bool(m.get("same_name"))})
            if similar:
                top = similar[0]
                warnings.append(("بنفس اسم نمط موجود" if top["same_name"] else f"بيشبه نمط موجود {round(top['score'] * 100)}%") + f": «{top['pattern_name']}»")
        rows.append({"row": p["row"], "name": name, "fields": fields,
                     "status": "error" if errs else "warn" if warnings else "ok",
                     "errors": errs, "warnings": warnings, "similar": similar})
    for key, scs in scenes_by_name.items():
        if key not in used_scene_keys:
            notes.append(f"مشاهد باسم نمط مش موجود بورقة الأنماط: «{key}» (سطر {scs[0]['_row']} بورقة المشاهد)")
    if not rows:
        raise ExcelError(f"ورقة «{PATTERNS_SHEET}» ما فيها أنماط — عبّي سطر على الأقل")

    _cleanup()
    token = secrets.token_hex(8)
    imports_dir().mkdir(parents=True, exist_ok=True)
    pattern_store.write_atomic(imports_dir() / f"{token}.json", {"created_at": time.time(), "filename": filename, "rows": rows})
    counts = {s: sum(1 for r in rows if r["status"] == s) for s in ("ok", "warn", "error")}
    return {"token": token, "filename": filename, "rows": [{k: v for k, v in r.items() if k != "fields"} for r in rows],
            "counts": counts, "notes": notes}


def _cleanup() -> None:
    d = imports_dir()
    if not d.exists():
        return
    now = time.time()
    for f in d.glob("*.json"):
        try:
            if now - f.stat().st_mtime > PREVIEW_TTL:
                f.unlink()
        except OSError:
            pass


# ═════════════════ التنفيذ ═════════════════
def commit(token: str, *, include_warnings: bool = False, rows: list[int] | None = None, by: str = "owner") -> dict:
    """يحفظ الأنماط الصح من معاينة (والمشابهة كمان إذا include_warnings). عناقيد مرة وحدة بالآخر."""
    if not TOKEN_RX.match(str(token or "")):
        raise ExcelError("المعاينة مش صالحة — ارفع الملف من جديد")
    f = imports_dir() / f"{token}.json"
    if not f.exists():
        raise ExcelError("المعاينة انتهت أو انستعملت — ارفع الملف من جديد")
    saved = json.loads(f.read_text(encoding="utf-8"))
    created, failed, skipped = [], [], 0
    for r in saved["rows"]:
        if rows is not None and r["row"] not in rows:
            skipped += 1
            continue
        if r["status"] == "error" or (r["status"] == "warn" and not include_warnings):
            skipped += 1
            continue
        try:
            pid = pattern_store.create(r["fields"], via="excel", by=by, extra_source={"file": saved.get("filename") or "", "row": r["row"]}, recluster=False)
            created.append({"row": r["row"], "id": pid, "name": r["name"]})
        except pattern_store.PatternError as e:
            failed.append({"row": r["row"], "name": r["name"], "error": str(e) if not isinstance(e.detail, list) else "، ".join(x["msg"] for x in e.detail)})
            if e.status == 423:                                      # المكتبة مشغولة: ما منكمّل نحاول على الباقي
                break
    if created:
        pattern_store.recluster()
        import pattern_i18n
        pattern_i18n.kick()                                      # لو في أنماط نصها بالإنجليزي بتنترجم لحالها
        f.unlink(missing_ok=True)                                # انستعملت. لو ما انحفظ شي بتضل لتعيد المحاولة
    return {"created": created, "failed": failed, "skipped": skipped}
